"""#342 production check: buffer cones on www.conestruct.com (via the site proxy).

Replays render342.py's two plans (flagger 45 mph rural, 500 ft; near side
25 mph urban, 500 ft) against production and checks that the audit, the
crew narrative, the XLSX and page 1 all carry the 3 buffer cones.

    python validation-artifacts/committed/issue-342-buffer-devices/prod-check/prodcheck342.py <repo-root> <out-dir>
"""

import io, json, os, re, sys, urllib.request

import pypdfium2 as pdfium
from openpyxl import load_workbook

root, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
tok = next(
    l.split("=", 1)[1].strip().strip('"')
    for l in open(os.path.join(root, "conestruct/site/.env.local"), encoding="utf-8")
    if l.startswith("GATE_BYPASS_TOKEN=")
)
base = "https://www.conestruct.com/api/render/"


def post(kind, body):
    req = urllib.request.Request(
        base + kind,
        data=json.dumps({"scenario": body}).encode(),
        method="POST",
        headers={"content-type": "application/json", "x-conestruct-gate": tok},
    )
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            return r.read()
    except urllib.error.HTTPError as e:
        raise SystemExit(f"{kind} {e.code}: {e.read()[:400]!r}") from e


META = {"project": "#342 prod check", "address": "", "lat": 0.0, "lng": 0.0, "siteConditions": {}}
COMMON = {"meta": META, "workType": "utility_cut", "duration": "short", "workLen": 500.0, "night": False, "divided": False}
# (name, body, buffer cones, spacing ft, page-1 label, total cones)
CASES = [
    (
        "flagger-45-rural-500",
        {**COMMON, "kind": "flagger_lane_closure", "roadType": "rural_undivided", "speed": 45, "lanes": 1,
         "laneWidth": 11.0, "pilotCar": False, "afad": False, "pedestrianAccess": False},
        3, 90.0, "BUFFER = 360 ft (NTS) · 3 CONES", 14,
    ),
    (
        "near-side-25-urban-500",
        {**COMMON, "kind": "near_intersection", "roadType": "urban_arterial", "speed": 25, "lanes": 2,
         "laneWidth": 12.0, "approaches": [{"id": "cross_n", "speed": 30, "roadType": "urban_arterial",
         "lanesPerDirection": 1, "laneWidth": 12.0, "signalized": False, "bearingDeg": None,
         "alongStationFt": -60.0}]},
        3, 38.8, "BUFFER = 155 ft (NTS) · 3 CONES", 15,
    ),
]
lines, ok = [], True
for name, body, n_buf, gap, label, total in CASES:
    sp = json.loads(post("audit", body))["sections"]["spacing"]
    md = post("markdown", body).decode("utf-8")
    xl = load_workbook(io.BytesIO(post("xlsx", body)))
    cones_xlsx = sum(r[5] for r in xl["Device List"].iter_rows(min_row=2, values_only=True) if r[2] and "Cone" in str(r[2]))
    doc = pdfium.PdfDocument(post("pdf", body))
    doc[0].render(scale=150 / 72).to_pil().convert("RGB").save(os.path.join(out, f"page1-{name}.png"), optimize=True)
    text = doc[0].get_textpage().get_text_range()
    summary = re.search(r"Traffic Cone \(\d+-inch\) (\d+)", text)
    step = re.search(r"Then continue the line through the buffer space:[^\n]*", md)
    checks = {
        "audit n_buffer_cones": sp.get("n_buffer_cones") == n_buf,
        "audit spacing": abs((sp.get("buffer_cone_spacing_ft") or 0) - gap) < 0.06,
        "audit total": sp.get("n_cones_placed_total") == total,
        "xlsx cones": cones_xlsx == total,
        "page 1 label": label in text,
        "page 1 summary": bool(summary) and int(summary.group(1)) == total,
        "crew step": bool(step) and f"{n_buf} more" in step.group(0),
    }
    passed = all(checks.values())
    ok &= passed
    lines.append(f"== {name}: {'PASS' if passed else 'FAIL'}")
    lines += [f"  {'ok ' if v else 'BAD'} {k}" for k, v in checks.items()]
    lines.append(f"  audit: {sp.get('buffer_count_text')}")
    lines.append(f"  totals: audit {sp.get('n_cones_placed_total')} | xlsx {cones_xlsx} | page 1 {summary.group(1) if summary else None}")
    lines.append(f"  crew: {step.group(0)[:160] if step else None}")
lines.append(f"RESULT: {'PASS' if ok else 'FAIL'} {sum(l.endswith('PASS') for l in lines if l.startswith('=='))}/{len(CASES)}")
open(os.path.join(out, "replay.txt"), "w", encoding="utf-8").write("\n".join(lines) + "\n")
print("\n".join(lines))
