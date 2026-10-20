"""#342 evidence: page 1, the audit's spacing lines and crew step 3 for the
issue's two live plans (flagger 45 mph rural, 500 ft zone; near side 25 mph
urban, 500 ft zone), rendered locally with the network stubbed.  Writes
<name>.png beside this file (or, with --out DIR, into DIR) and prints the
counts.  The divided lane closure is gated, so it has no API render; the
generator tests cover it.

Run from a checkout's root:
    python -B validation-artifacts/committed/issue-342-buffer-devices/render342.py [--out DIR]
"""

from __future__ import annotations

import io
import os
import re
import sys
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, os.getcwd())
os.environ.setdefault("RENDER_API_SECRET", "test-secret-do-not-deploy")

import pypdfium2 as pdfium  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from src.api import render_api  # noqa: E402
from src.rules import site_detection as sd  # noqa: E402

sd._overpass_request_with_fallback = lambda *a, **k: ({"elements": []}, None)  # type: ignore[assignment]
HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[sys.argv.index("--out") + 1]) if "--out" in sys.argv else HERE
OUT.mkdir(parents=True, exist_ok=True)
AUTH = {"Authorization": f"Bearer {os.environ['RENDER_API_SECRET']}"}
META = {"project": "#342", "address": "", "lat": 0.0, "lng": 0.0, "siteConditions": {}}

BODIES = {
    "flagger-45-rural-500": {
        "kind": "flagger_lane_closure",
        "meta": META,
        "roadType": "rural_undivided",
        "speed": 45,
        "lanes": 1,
        "laneWidth": 11.0,
        "divided": False,
        "workType": "utility_cut",
        "duration": "short",
        "workLen": 500.0,
        "night": False,
        "pilotCar": False,
        "afad": False,
        "pedestrianAccess": False,
    },
    "near-side-25-urban-500": {
        "kind": "near_intersection",
        "meta": META,
        "roadType": "urban_arterial",
        "speed": 25,
        "lanes": 2,
        "laneWidth": 12.0,
        "divided": False,
        "workType": "utility_cut",
        "duration": "short",
        "workLen": 500.0,
        "night": False,
        "approaches": [
            {
                "id": "cross_n",
                "speed": 30,
                "roadType": "urban_arterial",
                "lanesPerDirection": 1,
                "laneWidth": 12.0,
                "signalized": False,
                "bearingDeg": None,
                "alongStationFt": -60.0,
            }
        ],
    },
}

c = TestClient(render_api.app)
for name, b in BODIES.items():
    pdf = c.post("/render/pdf", headers=AUTH, json=b)
    pdf.raise_for_status()
    doc = pdfium.PdfDocument(pdf.content)
    doc[0].render(scale=150 / 72).to_pil().convert("RGB").save(OUT / f"{name}.png", optimize=True)
    text = doc[0].get_textpage().get_text_range()
    sp = c.post("/render/audit", headers=AUTH, json=b).json()["sections"]["spacing"]
    md = c.post("/render/markdown", headers=AUTH, json=b).text
    xl = load_workbook(io.BytesIO(c.post("/render/xlsx", headers=AUTH, json=b).content))
    cones_xlsx = sum(
        r[5]
        for r in xl["Device List"].iter_rows(min_row=2, values_only=True)
        if r[2] and "Cone" in str(r[2])
    )
    label = re.search(r"BUFFER = [^\n]*", text)
    summary = re.search(r"Traffic Cone \(\d+-inch\) (\d+)", text)
    step = re.search(r"Then continue the line through the buffer space:[^\n]*", md)
    print(f"== {name}")
    print("  page 1 label:", label.group(0).strip() if label else None)
    print(
        "  page 1 summary cones:",
        summary.group(1) if summary else None,
        "| xlsx cones:",
        cones_xlsx,
        "| audit placed total:",
        sp.get("n_cones_placed_total"),
    )
    print("  audit:", sp.get("buffer_count_text"))
    print("  audit:", sp.get("downstream_count_text"))
    print("  crew:", step.group(0)[:200] if step else None)
