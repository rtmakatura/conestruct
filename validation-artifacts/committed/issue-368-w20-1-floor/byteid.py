"""#368 byte-identity: hash every surface over every freeway shoulder plan
(posted 55-75, work-zone speed 20 to posted-5 and none, divided and
undivided, 1,000 and 5,000 ft) plus #362's non-freeway reduced-speed grid.

Run from the worktree root: PYTHONPATH=. uv run --extra dev python <this> <out.json>
Audit JSON (canonical), audit PDF text, plan PDF text, crew markdown and
XLSX cell values per plan (timestamps masked).  Compare a base run with an after run.
"""
import datetime, hashlib, io, itertools, json, sys
from openpyxl import load_workbook
from tests.s630._harness import _common
from tests.s630 import test_cross_surface as tcs

import os
from fastapi.testclient import TestClient
os.environ["RENDER_API_SECRET"] = "test-secret-do-not-deploy"
os.environ["MAPBOX_TOKEN"] = "test-token"
tcs._ps.httpx.get = lambda *a, **k: tcs._FakeTile()  # the fixture's offline tile
tcs._ps._validate_corridor_bearing = lambda corridor: None
api = TestClient(tcs._render_app)
H = tcs._API_HEADERS
h = lambda b: hashlib.sha256(b).hexdigest()[:16]
out = {}
FREEWAY = [(s, wz) for s in (55, 60, 65, 70, 75) for wz in [*range(20, s, 5), None]]
NONFWY = [(45, 35), (55, 45), (55, 50), (65, 55), (65, 60), (75, 65), (75, 70), (55, None),
          (55, 35), (65, 45), (65, 40), (75, 55), (75, 45), (75, 40)]
GRID = ([("freeway", div, p, wl) for div in (True, False) for p in FREEWAY for wl in (1000, 5000)]
        + [(road, road == "rural_divided", p, wl)
           for road in ("urban_arterial", "rural_undivided", "rural_divided")
           for p in NONFWY for wl in (200, 1000, 5000)])
for road, div, (s, wz), wl in GRID:
    body = {**_common(s, wz), "roadType": road, "divided": div, "workLen": wl}
    key = f"{road}{'' if road != 'freeway' else (' div' if div else ' undiv')} {s}->{wz} L{wl}"
    a = api.post("/render/audit", json=body, headers=H)
    if a.status_code != 200:
        out[key] = f"HTTP {a.status_code}"; continue
    aj = a.json()
    row = {"audit": h(json.dumps(aj, sort_keys=True).encode()),
           "note": any(i["kind"] == "fines_double_spacing"
                       for i in (aj.get("pending_verification") or {}).get("items", []))}
    ap = api.post("/render/audit-pdf", json=body, headers=H)
    row["audit_pdf"] = h("\n".join(tcs._pdf_text(ap.content)).encode()) if ap.status_code == 200 else ap.status_code
    pp = api.post("/render/pdf", json=body, headers=H)
    row["plan_pdf"] = h("\n".join(tcs._pdf_text(pp.content)).encode()) if pp.status_code == 200 else pp.status_code
    md = api.post("/render/markdown", json=body, headers=H)
    row["crew"] = h(md.content)
    xl = api.post("/render/xlsx", json=body, headers=H)
    if xl.status_code == 200:
        wb = load_workbook(io.BytesIO(xl.content))
        # The export timestamp (Summary!B8) is masked: it changes every run.
        cells = [[None if isinstance(c.value, datetime.datetime) else c.value for c in r]
                 for ws in wb for r in ws.iter_rows()]
        row["xlsx"] = h(repr(cells).encode())
    out[key] = row
json.dump(out, open(sys.argv[1], "w"), indent=1, sort_keys=True)
print(len(out), "plans;", sum(1 for v in out.values() if isinstance(v, dict) and v["note"]), "with the note")
