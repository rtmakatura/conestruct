import sys, os, io, copy, hashlib, json
root = sys.argv[1]; sys.path.insert(0, root); sys.path.insert(1, os.getcwd())
os.environ["RENDER_API_SECRET"] = "s"
from fastapi.testclient import TestClient
from src.api.render_api import app
from tests.test_plan_sheet_device_summary import FLAGGER_BODY, SHOULDER_BODY
from tests.test_r117_jurisdiction_raised import _broadway_denver
import src.api.render_api as r; assert os.path.normcase(r.__file__).startswith(os.path.normcase(os.path.normpath(root))), r.__file__
from openpyxl import load_workbook
c = TestClient(app); H = {"Authorization": "Bearer s"}
bodies = {}
for k, cls in [(None,None),("denver",None),("englewood","collector"),("greeley","arterial"),("lakewood","arterial"),("thornton",None),("parker",None)]:
    b = copy.deepcopy(FLAGGER_BODY)
    if k: b["jurisdiction_key"]=k
    if cls: b["street_class"]=cls
    bodies[f"flagger-{k}"]=b
    s = copy.deepcopy(SHOULDER_BODY)
    if k: s["jurisdiction_key"]=k
    if cls: s["street_class"]=cls
    bodies[f"shoulder-{k}"]=s
bodies["ni-denver"]=_broadway_denver()
def xl(content):
    wb=load_workbook(io.BytesIO(content)); out=[[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb]; wb.close(); import re; return re.sub(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?(\.\d+)?", "TS", json.dumps(out, default=str))
out={}
for name,b in bodies.items():
    row={}
    row["md"]=c.post("/render/markdown",json=b,headers=H).text
    row["xlsx"]=xl(c.post("/render/xlsx",json=b,headers=H).content)
    row["quote"]=xl(c.post("/render/quote",json={"scenario":b},headers=H).content)
    row["breakdown"]=c.post("/render/device-breakdown",json=b,headers=H).text
    row["audit"]=c.post("/render/audit",json=b,headers=H).text
    out[name]={k:hashlib.sha256(v.encode()).hexdigest()[:12] for k,v in row.items()}
print(json.dumps(out))
