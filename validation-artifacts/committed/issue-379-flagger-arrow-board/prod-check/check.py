"""#379 production check: every Denver flagger row says caution mode.

Run from the repo root: ``uv run python validation-artifacts/committed/issue-379-flagger-arrow-board/prod-check/check.py``.
It replays the audit lane's own requests (#370, ``audit-lane/runs/<id>/request.json``)
through www.conestruct.com's ``/api/render/*`` proxy (the lane's gate headers):
the audit with ``include_breakdown`` and a scan token, then the five downloads
carrying that token.  The token stays in memory and is never written (#371).
Writes ``results.json`` beside itself.

Rows: F1-F4, F6, F18 (the six Denver flagger rows, #379 acceptance) and N1
(Denver near-intersection, the unchanged control).
"""

from __future__ import annotations

import io
import json
import re
import sys
import time
from datetime import UTC, datetime
from pathlib import Path

import httpx
import pypdfium2 as pdfium
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
sys.path.insert(0, str(REPO / "scripts"))
from gate import gate_headers  # noqa: E402

BASE = "https://www.conestruct.com"
RUNS = REPO / "validation-artifacts/committed/audit-lane/runs"
FLAGGER = ["F1", "F2", "F3", "F4", "F6", "F18"]
CONTROL = ["N1"]
NAME = "Arrow Board (caution mode only)"
CITE = "§6L.06 ¶18"
GAP_S = 4.0


class Paced:
    def __init__(self) -> None:
        self.c = httpx.Client(
            headers={**gate_headers(BASE), "content-type": "application/json"}, timeout=300
        )
        self.last = 0.0

    def post(self, path: str, body: dict) -> httpx.Response:
        for attempt in range(4):
            wait = self.last + GAP_S - time.time()
            if wait > 0:
                time.sleep(wait)
            r = self.c.post(f"{BASE}{path}", content=json.dumps(body))
            self.last = time.time()
            print(time.strftime("%H:%M:%S"), path, r.status_code, flush=True)
            if r.status_code == 429:
                time.sleep(65)
                continue
            if r.status_code == 400 and "site_scan_unavailable" in r.text[:2000]:
                time.sleep(30 * 2**attempt)
                continue
            return r
        return r


def page1(content: bytes) -> str:
    pdf = pdfium.PdfDocument(content)
    try:
        return pdf[0].get_textpage().get_text_range()
    finally:
        pdf.close()


def all_pages(content: bytes) -> str:
    pdf = pdfium.PdfDocument(content)
    try:
        return "\n".join(p.get_textpage().get_text_range() for p in pdf)
    finally:
        pdf.close()


def xlsx_rows(content: bytes, sheet: str) -> list[list]:
    wb = load_workbook(io.BytesIO(content))
    try:
        return [list(r) for r in wb[sheet].iter_rows(min_row=2, values_only=True)]
    finally:
        wb.close()


def arrow(rows: list[list], col: int) -> list[list]:
    return [r for r in rows if r and "Arrow" in str(r[col])]


def check_row(p: Paced, rid: str) -> dict:
    sc = json.loads((RUNS / rid / "request.json").read_text("utf-8"))["scenario"]
    out: dict = {"kind": sc["kind"], "jurisdiction_key": sc.get("jurisdiction_key")}
    a = p.post(
        "/api/render/audit",
        {"scenario": {**sc, "include_scan_token": True, "include_breakdown": True}},
    )
    out["audit_status"] = a.status_code
    if a.status_code != 200:
        out["audit_error"] = a.text[:500]
        return out
    body = a.json()
    token = body.pop("site_scan_token", None)
    bd = body.get("breakdown") or {}
    out["audit"] = {
        "breakdown_arrow_rows": [
            {"device": d["device"], "qty": d["qty"]}
            for d in bd.get("devices", [])
            if "Arrow" in d["device"]
        ],
        "arrow_rules": [
            {k: d.get(k) for k in ("device_label", "mode_note", "raised")}
            for d in (bd.get("jurisdiction") or {}).get("applied_deltas", [])
            if "arrow" in str((d.get("effect") or {}).get("device"))
        ],
    }
    dl = {**sc}
    if token:
        dl["site_scan"] = {"proceed_if_unavailable": False, "token": token}
    got: dict[str, bytes] = {}
    for kind in ("pdf", "xlsx", "markdown", "crew-pdf", "quote"):
        r = p.post(f"/api/render/{kind}", {"scenario": dl})
        out[f"{kind}_status"] = r.status_code
        if r.status_code == 200:
            got[kind] = r.content
    if "pdf" in got:
        t = page1(got["pdf"])
        out["page1_arrow_lines"] = [ln.strip() for ln in t.splitlines() if "Arrow" in ln]
    if "xlsx" in got:
        out["xlsx_arrow_rows"] = [
            {"name": r[1], "qty": r[5], "note": r[6]} for r in arrow(xlsx_rows(got["xlsx"], "Device List"), 1)
        ]
    if "quote" in got:
        out["quote_arrow_rows"] = [
            {"name": r[1], "qty": r[4], "note": r[9]}
            for r in arrow(xlsx_rows(got["quote"], "Equipment Detail"), 1)
        ]
    if "markdown" in got:
        md = got["markdown"].decode("utf-8")
        out["crew_md"] = {
            "equipment_line": next(
                (ln for ln in md.splitlines() if "Arrow Board" in ln and ln.startswith("- ")), None
            ),
            "setup_paragraph": next(
                (ln.strip() for ln in md.splitlines() if ln.strip().startswith("Arrow board (")), None
            ),
            "takedown_line": next((ln for ln in md.splitlines() if "remove the arrow board" in ln), None),
            "states_arrow_mode": bool(re.search(r"ARROW mode|chevron mode", md)),
        }
    if "crew-pdf" in got:
        t = all_pages(got["crew-pdf"])
        out["crew_pdf_says_caution"] = "caution mode only" in t
    return out


def verdict(rid: str, r: dict) -> list[str]:
    """Failures for one row; empty is PASS."""
    bad: list[str] = []
    statuses = [r.get(k) for k in ("audit_status", "pdf_status", "xlsx_status", "markdown_status",
                                   "crew-pdf_status", "quote_status")]
    if any(s != 200 for s in statuses):
        return [f"status {statuses}"]
    if rid in FLAGGER:
        if not any(NAME in ln for ln in r["page1_arrow_lines"]):
            bad.append("page 1")
        if [x["name"] for x in r["xlsx_arrow_rows"]] != [NAME] or CITE not in r["xlsx_arrow_rows"][0]["note"]:
            bad.append("device xlsx")
        if [x["name"] for x in r["quote_arrow_rows"]] != [NAME] or CITE not in r["quote_arrow_rows"][0]["note"]:
            bad.append("quote xlsx")
        crew = r["crew_md"]
        if not (crew["equipment_line"] and NAME in crew["equipment_line"]):
            bad.append("crew equipment")
        if not (crew["setup_paragraph"] and CITE in crew["setup_paragraph"]):
            bad.append("crew setup")
        if not crew["takedown_line"] or crew["states_arrow_mode"] or not r["crew_pdf_says_caution"]:
            bad.append("crew takedown/mode/pdf")
        if [x["device"] for x in r["audit"]["breakdown_arrow_rows"]] != [NAME]:
            bad.append("audit breakdown")
        rules = r["audit"]["arrow_rules"]
        if not rules or not all(CITE in (x["mode_note"] or "") for x in rules):
            bad.append("audit rule")
    else:
        if any("caution mode only" in str(x) for x in (r["xlsx_arrow_rows"], r["quote_arrow_rows"],
                                                         r["audit"], r["page1_arrow_lines"])):
            bad.append("control changed")
        if r["crew_md"]["setup_paragraph"] is not None:
            bad.append("control crew changed")
    return bad


def main() -> None:
    p = Paced()
    healthz = httpx.get(
        "https://rtmakatura--conestruct-render-fastapi-app.modal.run/healthz", timeout=60
    ).json()
    results: dict = {
        "base": BASE,
        "checked_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "healthz_sha": healthz.get("sha") or healthz,
        "rows": {},
    }
    for rid in FLAGGER + CONTROL:
        r = check_row(p, rid)
        r["fail"] = verdict(rid, r)
        r["verdict"] = "FAIL" if r["fail"] else "PASS"
        results["rows"][rid] = r
        print(rid, r["verdict"], r["fail"], flush=True)
    passed = sum(r["verdict"] == "PASS" for r in results["rows"].values())
    results["summary"] = f"{passed}/{len(results['rows'])} PASS"
    (HERE / "results.json").write_text(json.dumps(results, indent=1, ensure_ascii=False) + "\n", "utf-8")
    print(results["summary"])


if __name__ == "__main__":
    main()
