"""#374 production check: a flagger plan's AFADs carry one name on every surface.

Run from the repo root with GATE_BYPASS_TOKEN set:
``uv run python validation-artifacts/committed/issue-374-afad-name/prod-check/check.py``.
It replays the audit lane's own requests (#370, ``audit-lane/runs/<id>/request.json``)
through www.conestruct.com's ``/api/render/*`` proxy: the audit (with its
breakdown) for a scan token, then page 1 (summary and legend), the device XLSX,
the quote XLSX, the quote breakdown, the device breakdown and the crew sheet
carrying it.  The token stays in memory and is never written (#371).  Writes
``results.json``.

Rows: F16 (flagger, ``afad: true``, the #374 acceptance: the AFAD name, qty 2,
on every surface, no old name) and F1 (flagger, no AFAD, the control: no AFAD
name, no temporary signal anywhere).
"""

from __future__ import annotations

import io
import json
import sys
import time
from datetime import UTC, datetime
from pathlib import Path

import httpx
import pypdfium2 as pdfium
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
sys.path.insert(0, str(REPO / "scripts"))
from gate import gate_headers  # noqa: E402

BASE = "https://www.conestruct.com"
RUNS = REPO / "validation-artifacts/committed/audit-lane/runs"
NAME = "Automated Flagger Assistance Device (AFAD)"
OLD_NAMES = ("Portable traffic signal", "Temporary Traffic Signal", "Temporary Signal")
GAP_S = 4.0


class Paced:
    def __init__(self) -> None:
        self.c = httpx.Client(
            headers={**gate_headers(BASE), "content-type": "application/json"}, timeout=300
        )
        self.last = 0.0

    def post(self, path: str, body: dict) -> httpx.Response:
        for attempt in range(4):
            wait = self.last + GAP_S - time.time()
            if wait > 0:
                time.sleep(wait)
            r = self.c.post(f"{BASE}{path}", content=json.dumps(body))
            self.last = time.time()
            print(time.strftime("%H:%M:%S"), path, r.status_code, flush=True)
            if r.status_code == 429:
                time.sleep(65)
                continue
            if r.status_code == 400 and "site_scan_unavailable" in r.text[:2000]:
                time.sleep(30 * 2**attempt)
                continue
            return r
        return r


def _old(texts: list[str]) -> list[str]:
    return [t for t in texts if any(o in t for o in OLD_NAMES)]


def page1(content: bytes) -> dict:
    pdf = pdfium.PdfDocument(content)
    try:
        pages = [p.get_textpage().get_text_range().splitlines() for p in pdf]
    finally:
        pdf.close()
    lines = pages[0]
    s = lines.index("TRAFFIC CONTROL DEVICE SUMMARY") + 2
    e = next(i for i in range(s, len(lines)) if lines[i].startswith("TOTAL DEVICES"))
    rows = [ln.removeprefix("— ").strip() for ln in lines[s:e]]
    return {
        "summary_afad_rows": [r for r in rows if r.startswith(NAME)],
        # The legend row is the name alone on its line (no quantity).
        "legend_afad_lines": [ln for ln in lines[:s] if ln.strip() == NAME],
        "old_name_lines": _old([ln for page in pages for ln in page]),
    }


def _xlsx_rows(content: bytes, sheet: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(content))
    try:
        return list(wb[sheet].iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()


def device_xlsx(content: bytes) -> dict:
    rows = _xlsx_rows(content, "Device List")
    return {
        "signal_rows": [list(r[1:]) for r in rows if r[1] == "TEMPORARY_SIGNAL"],
        "old_name_cells": _old([str(c) for r in rows for c in r]),
    }


def quote_xlsx(content: bytes) -> dict:
    rows = _xlsx_rows(content, "Equipment Detail")
    return {
        "signal_rows": [list(r[1:]) for r in rows if r[1] == "TEMPORARY_SIGNAL"],
        "old_name_cells": _old([str(c) for r in rows for c in r]),
    }


def check_row(p: Paced, rid: str) -> dict:
    sc = json.loads((RUNS / rid / "request.json").read_text("utf-8"))["scenario"]
    out: dict = {"kind": sc["kind"], "afad": sc.get("afad")}
    a = p.post(
        "/api/render/audit",
        {"scenario": {**sc, "include_scan_token": True, "include_breakdown": True}},
    )
    out["audit_status"] = a.status_code
    if a.status_code != 200:
        out["audit_error"] = a.text[:500]
        return out
    audit = a.json()
    token = audit.get("site_scan_token")
    devices = (audit.get("breakdown") or {}).get("devices", [])
    out["audit_breakdown_afad"] = [d for d in devices if d["device"] == NAME]
    out["audit_breakdown_old"] = _old([d["device"] for d in devices])
    out["audit_stations_text"] = (audit.get("sections", {}).get("flagger") or {}).get(
        "stations_text"
    )
    dl = {**sc}
    if token:
        dl["site_scan"] = {"proceed_if_unavailable": False, "token": token}
    body = {"scenario": dl}
    for key, path, parse in (
        ("page1", "/api/render/pdf", page1),
        ("device_xlsx", "/api/render/xlsx", device_xlsx),
        ("quote_xlsx", "/api/render/quote", quote_xlsx),
    ):
        r = p.post(path, body)
        out[f"{key}_status"] = r.status_code
        if r.status_code == 200:
            out[key] = parse(r.content)
    r = p.post("/api/render/quote-breakdown", body)
    out["quote_breakdown_status"] = r.status_code
    if r.status_code == 200:
        lines = r.json()["equipment_lines"]
        out["quote_breakdown"] = {
            "signal_lines": [
                {k: ln[k] for k in ("description", "qty", "unit", "daily_rate", "extended")}
                for ln in lines
                if ln["device_type"] == "TEMPORARY_SIGNAL"
            ],
            "old_names": _old([ln["description"] for ln in lines]),
        }
    r = p.post("/api/render/device-breakdown", body)
    out["breakdown_status"] = r.status_code
    if r.status_code == 200:
        devs = r.json()["devices"]
        out["breakdown"] = {
            "afad_rows": [d for d in devs if d["device"] == NAME],
            "old_names": _old([d["device"] for d in devs]),
        }
    r = p.post("/api/render/markdown", body)
    out["crew_status"] = r.status_code
    if r.status_code == 200:
        md = r.text.splitlines()
        out["crew"] = {
            "afad_lines": [ln.strip() for ln in md if NAME in ln],
            "old_name_lines": _old(md),
        }
    return out


def verdict(r: dict) -> list[str]:
    keys = [k for k in r if k.endswith("_status")]
    if any(r[k] != 200 for k in keys):
        return [f"status {[(k, r[k]) for k in keys if r[k] != 200]}"]
    bad = []
    old = {
        "page1": r["page1"]["old_name_lines"],
        "device_xlsx": r["device_xlsx"]["old_name_cells"],
        "quote_xlsx": r["quote_xlsx"]["old_name_cells"],
        "quote_breakdown": r["quote_breakdown"]["old_names"],
        "breakdown": r["breakdown"]["old_names"],
        "audit_breakdown": r["audit_breakdown_old"],
        "crew": r["crew"]["old_name_lines"],
    }
    bad += [f"{k} still carries an old name: {v}" for k, v in old.items() if v]
    if r["afad"]:
        if r["page1"]["summary_afad_rows"] != [f"{NAME} 2"]:
            bad.append(f"page 1 summary {r['page1']['summary_afad_rows']}")
        if len(r["page1"]["legend_afad_lines"]) != 1:
            bad.append(f"page 1 legend {r['page1']['legend_afad_lines']}")
        x = r["device_xlsx"]["signal_rows"]
        if len(x) != 1 or x[0][1:5] != [NAME, "630-86810", "EACH", 2]:
            bad.append(f"device XLSX {x}")
        q = r["quote_xlsx"]["signal_rows"]
        if len(q) != 1 or (q[0][2], q[0][3]) != (NAME, 2):
            bad.append(f"quote XLSX {q}")
        ql = r["quote_breakdown"]["signal_lines"]
        if [(ln["description"], ln["qty"]) for ln in ql] != [(NAME, 2)]:
            bad.append(f"quote breakdown {ql}")
        want = [{"device": NAME, "code": "—", "function": "Traffic control", "qty": 2}]
        if r["breakdown"]["afad_rows"] != want:
            bad.append(f"breakdown {r['breakdown']['afad_rows']}")
        if r["audit_breakdown_afad"] != want:
            bad.append(f"audit breakdown {r['audit_breakdown_afad']}")
        if r["audit_stations_text"] != "2 AFAD":
            bad.append(f"audit stations_text {r['audit_stations_text']!r}")
        if r["crew"]["afad_lines"] != [f"- 2× {NAME}"]:
            bad.append(f"crew {r['crew']['afad_lines']}")
    else:
        hits = (
            r["page1"]["summary_afad_rows"]
            + r["page1"]["legend_afad_lines"]
            + r["device_xlsx"]["signal_rows"]
            + r["quote_xlsx"]["signal_rows"]
            + r["quote_breakdown"]["signal_lines"]
            + r["breakdown"]["afad_rows"]
            + r["audit_breakdown_afad"]
            + r["crew"]["afad_lines"]
        )
        if hits:
            bad.append(f"control lists an AFAD or temporary signal: {hits}")
    return bad


def main() -> None:
    p = Paced()
    healthz = httpx.get(
        "https://rtmakatura--conestruct-render-fastapi-app.modal.run/healthz", timeout=60
    ).json()
    results: dict = {
        "base": BASE,
        "checked_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "healthz_sha": healthz.get("sha") or healthz,
        "rows": {},
    }
    for rid in ("F16", "F1"):
        r = check_row(p, rid)
        r["fail"] = verdict(r)
        r["verdict"] = "FAIL" if r["fail"] else "PASS"
        results["rows"][rid] = r
        print(rid, r["verdict"], r["fail"], flush=True)
    passed = sum(r["verdict"] == "PASS" for r in results["rows"].values())
    results["summary"] = f"{passed}/{len(results['rows'])} PASS"
    (HERE / "results.json").write_text(
        json.dumps(results, indent=1, ensure_ascii=False) + "\n", "utf-8"
    )
    print(results["summary"])


if __name__ == "__main__":
    main()
