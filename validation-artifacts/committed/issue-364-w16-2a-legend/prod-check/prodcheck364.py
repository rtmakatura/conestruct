"""#364 production check: the W16-2a plaque legend on www.conestruct.com.

A Case 11 plan's W16-2a reads "2,000 FT" (R1); a Case 26/27 plan reads
Sheet 14's "¼ MILE" (R2).  No surface reads "W16-2a ... NEXT".
Checks the audit, page 1, the crew sheet and the device list, the same
strings as tests/s630/test_cross_surface.py's #364 test.
Usage: PYTHONPATH=<worktree> python prodcheck364.py <main checkout root>
"""
import io, json, os, re, sys, urllib.error, urllib.request

import pypdfium2 as pdfium
from openpyxl import load_workbook

from tests.s630.test_cross_surface import BODIES

root = sys.argv[1]
tok = next(
    l.split("=", 1)[1].strip().strip('"')
    for l in open(os.path.join(root, "conestruct/site/.env.local"), encoding="utf-8")
    if l.startswith("GATE_BYPASS_TOKEN=")
)
base = "https://www.conestruct.com/api/render/"


def post(kind, body):
    req = urllib.request.Request(
        base + kind,
        data=json.dumps({"scenario": body}).encode(),
        method="POST",
        headers={"content-type": "application/json", "x-conestruct-gate": tok},
    )
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def check(name, good, shown):
    print(f"   {'ok' if good else 'FAIL'} {name}: {ascii(shown)}")
    return good


PLANS = [
    ("Case 11 (55 -> 50)", BODIES["case_11"], "2,000 FT"),
    ("Case 26 (65 -> 60)", BODIES["case_26"], "¼ MILE"),
    ("Case 27", BODIES["case_27"], "¼ MILE"),
]

ok = True
for n, (label, body, legend) in enumerate(PLANS, 1):
    print(f"{n}. {label}, expect {ascii(legend)}")
    statuses = {}
    s, audit = post("audit", body)
    statuses["audit"] = s
    rows = [] if s != 200 else [
        r["Distance from Taper (ft)"]
        for r in json.loads(audit)["sections"]["advance"]["sign_table"]
        if r["Code"] == "W16-2a"
    ]
    ok &= check("audit", s == 200 and rows and all(
        r == f"{legend} (under first W21-5aR)" for r in rows), rows)

    s, pdf = post("pdf", body)
    statuses["pdf"] = s
    page1 = "" if s != 200 else " ".join(
        pdfium.PdfDocument(pdf)[0].get_textpage().get_text_range().split())
    hit = re.findall(r"W16-2a \S+ \S+", page1)
    ok &= check("page 1", s == 200 and f"W16-2a {legend}" in page1, hit)

    s, md = post("markdown", body)
    statuses["markdown"] = s
    md = md.decode("utf-8") if s == 200 else ""
    hit = [l for l in md.splitlines() if "| W16-2a |" in l]
    ok &= check("crew sheet", f"| W16-2a | {legend} (under upstream W21-5aR) |" in md, hit)

    s, xb = post("xlsx", body)
    statuses["xlsx"] = s
    cells = []
    if s == 200:
        wb = load_workbook(io.BytesIO(xb), read_only=True)
        cells = [str(c) for row in wb["Device List"].iter_rows(values_only=True)
                 for c in row if "W16-2a" in str(c)]
        wb.close()
    ok &= check("device list", cells == [f"W16-2a {legend}"], cells)

    nxt = any(re.search(r"W16-2a[\s|]*NEXT", t) for t in (page1, md, " ".join(cells)))
    ok &= check("no NEXT", not nxt, "absent" if not nxt else "PRESENT")
    print(f"   statuses: {statuses}")

print("PASS 3/3" if ok else "FAIL")
sys.exit(0 if ok else 1)
