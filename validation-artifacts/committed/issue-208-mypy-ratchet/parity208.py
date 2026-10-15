"""#208 parity: every deliverable route, every corpus scenario, before vs after.

Run from a checkout's root:  python -B parity208.py <out.json>
It posts every corpus case (tests/corpus/manifest.py), plus the
R117 near-intersection N Broadway bodies (two-way and one-way), through the
FastAPI app in-process (no network: no ``site_scan`` on these bodies), to every
route that builds a deliverable or a wire answer:

  /render/audit, /render/device-breakdown, /render/quote-breakdown (JSON)
  /render/xlsx, /render/quote (every cell value, every sheet)
  /render/markdown (text)
  /render/pdf, /render/crew-pdf, /render/audit-pdf (every page's text)

Render dates are masked (the DATE cell is today's date; the workbook's
"Generated" cell is the render's datetime). Each output is
hashed; the JSON written maps case -> route -> sha256. Two runs compare
equal iff every route answered the same status and the same content.
"""

from __future__ import annotations

import datetime
import hashlib
import io
import json
import os
import re
import sys
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, os.getcwd())
os.environ.setdefault("RENDER_API_SECRET", "test-secret-do-not-deploy")

import openpyxl  # noqa: E402
import pypdfium2 as pdfium  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import tests.test_r117_pdf_notes as r117  # noqa: E402
from src.api.render_api import app  # noqa: E402
from tests.corpus.manifest import CASES, scenario_body  # noqa: E402

AUTH = {"Authorization": "Bearer test-secret-do-not-deploy"}
_DATES = re.compile(
    r"\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4}|"
    r"(January|February|March|April|May|June|July|August|September|October|"
    r"November|December) \d{1,2}, \d{4}"
)


def _mask(text: str) -> str:
    return _DATES.sub("<date>", text)


def _pdf_text(blob: bytes) -> str:
    doc = pdfium.PdfDocument(blob)
    pages = [doc[i].get_textpage().get_text_range() for i in range(len(doc))]
    doc.close()
    return "\f".join(pages)


def _xlsx_text(blob: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            # The "Generated" cell is a datetime of the render; mask it.
            out.append(
                repr(
                    tuple("<datetime>" if isinstance(v, datetime.datetime) else v for v in row)
                )
            )
    return "\n".join(out)


ROUTES = {
    "/render/audit": "json",
    "/render/device-breakdown": "json",
    "/render/quote-breakdown": "json",
    "/render/xlsx": "xlsx",
    "/render/quote": "xlsx",
    "/render/markdown": "text",
    "/render/pdf": "pdf",
    "/render/crew-pdf": "pdf",
    "/render/audit-pdf": "pdf",
}

QUOTE_ROUTES = {"/render/quote", "/render/quote-breakdown"}


def _content(kind: str, r) -> str:
    if r.status_code != 200:
        return f"{r.status_code} {r.text}"
    if kind == "json":
        return json.dumps(r.json(), sort_keys=True)
    if kind == "xlsx":
        return _xlsx_text(r.content)
    if kind == "pdf":
        return _pdf_text(r.content)
    return r.text


def _no_network() -> None:
    """The corpus conftest's block (tests/corpus/conftest.py:44-80), and no
    Mapbox token, so neither side's output depends on the wire or on which
    checkout's ``.env`` ``load_dotenv`` found."""
    import httpx
    import requests

    from src.rules import site_detection

    def _blocked(*_a: object, **_k: object) -> None:
        raise RuntimeError("parity208: no network")

    async def _blocked_mirror(*_a: object, **_k: object) -> None:
        _blocked()

    for name in ("get", "post", "put", "patch", "delete", "head", "request", "stream"):
        setattr(httpx, name, _blocked)
    site_detection._mirror_post = _blocked_mirror  # type: ignore[assignment]
    for name in ("get", "post", "put", "patch", "delete", "head", "request"):
        setattr(requests, name, _blocked)
    requests.sessions.Session.request = _blocked  # type: ignore[method-assign,assignment]
    os.environ.pop("MAPBOX_TOKEN", None)


def main(out_path: str) -> None:
    _no_network()
    bodies = {c.id: scenario_body(c) for c in CASES}
    bodies["r117-broadway-twoway"] = dict(r117.BROADWAY)
    bodies["r117-broadway-oneway"] = {**r117.BROADWAY, "carriageway": r117.ONE_WAY}
    client = TestClient(app)
    result: dict[str, dict[str, str]] = {}
    for case_id, body in sorted(bodies.items()):
        result[case_id] = {}
        for route, kind in ROUTES.items():
            # The quote routes take {scenario, settings}; default settings.
            payload = {"scenario": body} if route in QUOTE_ROUTES else body
            r = client.post(route, headers=AUTH, json=payload)
            text = _mask(_content(kind, r))
            result[case_id][route] = f"{r.status_code} " + hashlib.sha256(text.encode()).hexdigest()
    Path(out_path).write_text(json.dumps(result, indent=1, sort_keys=True), encoding="utf-8")
    n = sum(len(v) for v in result.values())
    ok = sum(1 for v in result.values() for h in v.values() if h.startswith("200 "))
    print(f"{len(result)} scenarios, {n} responses, {ok} answered 200 -> {out_path}")


if __name__ == "__main__":
    main(sys.argv[1])
