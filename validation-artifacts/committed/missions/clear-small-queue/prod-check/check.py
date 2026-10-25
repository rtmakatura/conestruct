"""Mission clear-small-queue (#392) production check of bc3b7d4.

    #389 #391 #353  healthz sha == bc3b7d4 (test-only changes)
    #388  Thornton "Advance Warning Signs (jurisdiction-required)" count on the
          device XLSX, quote XLSX, device breakdown and crew sheet (markdown):
          two-way undivided shoulder 4, divided shoulder 2, one-way shoulder 2,
          flagger 4, near_intersection two-way 1 leg 6; every surface agrees.
          Control: the shoulder body with CDOT, Lakewood and no jurisdiction
          prints no such row.
    #390  POST /render/corridor-map answers 200 image/jpeg, FF D8 FF bytes.

Bodies are the ones tests/test_thornton_stipulation6_count_388.py and
issue-388/churn.py use (imported from the tests).  Requests go straight to the
Modal backend (Bearer MODAL_RENDER_SECRET from conestruct/site/.env.local,
never printed or saved).

    MODAL_RENDER_SECRET=... uv run --extra dev python \
        validation-artifacts/committed/missions/clear-small-queue/prod-check/check.py
"""

import copy
import io
import json
import os
import re
import sys
from pathlib import Path
from typing import Any

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(REPO))

import httpx  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from tests.test_plan_sheet_device_summary import FLAGGER_BODY, SHOULDER_BODY  # noqa: E402
from tests.test_r117_jurisdiction_raised import _broadway_denver  # noqa: E402

BACKEND = "https://rtmakatura--conestruct-render-fastapi-app.modal.run"
EXPECT_SHA = "bc3b7d4da50579d54c2d1c0c71ce0d9376a3e750"
AWS = "Advance Warning Signs"

# issue-388/churn.py SHOULDER_2WAY / SHOULDER_1WAY
ONE_WAY = {"oneway": "yes", "highwayClass": "residential", "twinSearched": True}
SHOULDER_2WAY = {**SHOULDER_BODY, "roadType": "rural_undivided", "divided": False, "lanes": 1}
SHOULDER_1WAY = {
    **SHOULDER_BODY,
    "roadType": "urban_arterial",
    "speed": 30,
    "divided": False,
    "carriageway": ONE_WAY,
}


def with_j(base: dict, key: str | None) -> dict:
    b = copy.deepcopy(base)
    b.pop("jurisdiction_key", None)
    if key:
        b["jurisdiction_key"] = key
    return b


def xlsx_row(content: bytes) -> list | None:
    wb = load_workbook(io.BytesIO(content))
    try:
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                for v in r:
                    if isinstance(v, str) and v.startswith(AWS):
                        return [x for x in r if x is not None]
        return None
    finally:
        wb.close()


rows: list[dict[str, Any]] = []


def row(check: str, case: str, ok: bool | None, observed: Any) -> None:
    rows.append({"check": check, "case": case, "pass": ok, "observed": observed})
    v = {True: "PASS", False: "FAIL", None: "N/A"}[ok]
    print(f"{v} | {check} | {case} | {str(observed)[:240]}")


def surfaces(api: httpx.Client, body: dict) -> dict[str, Any]:
    out: dict[str, Any] = {}
    x = api.post(f"{BACKEND}/render/xlsx", json=body)
    out["xlsx"] = {
        "status": x.status_code,
        "row": xlsx_row(x.content) if x.status_code == 200 else None,
    }
    q = api.post(f"{BACKEND}/render/quote", json={"scenario": body})
    out["quote"] = {
        "status": q.status_code,
        "row": xlsx_row(q.content) if q.status_code == 200 else None,
    }
    bd = api.post(f"{BACKEND}/render/device-breakdown", json=body)
    br = None
    if bd.status_code == 200:
        br = [d for d in bd.json()["devices"] if str(d.get("device", "")).startswith(AWS)] or None
    out["breakdown"] = {"status": bd.status_code, "row": br}
    md = api.post(f"{BACKEND}/render/markdown", json=body)
    m = re.findall(r"- (\d+)× (Advance Warning Signs[^\n]*)", md.text)
    out["crew_sheet"] = {"status": md.status_code, "row": m or None}
    return out


def counts(s: dict[str, Any]) -> dict[str, Any]:
    """The Advance Warning Signs quantity each surface prints (None = no row).

    XLSX rows: the first number after the label is the quantity.
    """
    c: dict[str, Any] = {}
    for k in ("xlsx", "quote"):
        r = s[k]["row"]
        nums = [v for v in r[1:] if isinstance(v, (int, float))] if r else []
        c[k] = nums[0] if nums else None
    br = s["breakdown"]["row"]
    c["breakdown"] = br[0]["qty"] if br else None
    md = s["crew_sheet"]["row"]
    c["crew_sheet"] = int(md[0][0]) if md else None
    return c


def main() -> int:
    sha = httpx.get(f"{BACKEND}/healthz", timeout=60).json().get("sha", "")
    row("#389 #391 #353", "healthz sha == bc3b7d4", sha == EXPECT_SHA, f"healthz {sha}")

    secret = os.environ["MODAL_RENDER_SECRET"]
    api = httpx.Client(headers={"authorization": f"Bearer {secret}"}, timeout=240)

    raw: dict[str, Any] = {}
    cases = (
        ("Thornton shoulder, two-way undivided", with_j(SHOULDER_2WAY, "thornton"), 4),
        ("Thornton shoulder, divided", with_j(SHOULDER_BODY, "thornton"), 2),
        ("Thornton shoulder, one-way", with_j(SHOULDER_1WAY, "thornton"), 2),
        ("Thornton flagger_lane_closure", with_j(FLAGGER_BODY, "thornton"), 4),
        ("Thornton near_intersection, two-way, 1 leg", with_j(_broadway_denver(), "thornton"), 6),
        ("CDOT shoulder, two-way undivided (control)", with_j(SHOULDER_2WAY, "cdot"), None),
        ("Lakewood shoulder, two-way undivided (control)", with_j(SHOULDER_2WAY, "lakewood"), None),
        (
            "no jurisdiction shoulder, two-way undivided (control)",
            with_j(SHOULDER_2WAY, None),
            None,
        ),
        ("CDOT flagger (control)", with_j(FLAGGER_BODY, "cdot"), None),
    )
    for name, body, want in cases:
        s = surfaces(api, body)
        raw[name] = s
        c = counts(s)
        statuses = {k: v["status"] for k, v in s.items()}
        ok = all(v == 200 for v in statuses.values()) and all(v == want for v in c.values())
        row("#388", f"{name} (expect {want})", ok, {"counts": c, "status": statuses})

    scen = json.loads(
        (REPO / "tests/fixtures/corridor/broadway-sb.json").read_text(encoding="utf-8")
    )
    pic = {"scenario": scen, "stage": "laid_out", "width": 600, "height": 300}
    r = api.post(f"{BACKEND}/render/corridor-map", json=pic)
    ct = r.headers.get("content-type", "")
    row(
        "#390",
        "backend /render/corridor-map broadway-sb laid_out 600x300",
        r.status_code == 200 and ct.startswith("image/jpeg") and r.content[:3] == b"\xff\xd8\xff",
        f"{r.status_code}; content-type {ct}; first bytes {r.content[:3].hex(' ').upper()}; "
        f"{len(r.content)} bytes",
    )
    row(
        "#390",
        "proxy: 200 with no content-type answers 502",
        None,
        "unreachable in production (the backend always sends a content-type); tests only",
    )

    (HERE / "results.json").write_text(
        json.dumps({"healthz_sha": sha, "rows": rows}, indent=1, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    (HERE / "surfaces-raw.json").write_text(
        json.dumps(raw, indent=1, ensure_ascii=False, default=str) + "\n", encoding="utf-8"
    )
    fails = [r for r in rows if r["pass"] is False]
    print(
        f"healthz {sha[:7]}; {sum(r['pass'] is True for r in rows)} PASS, "
        f"{len(fails)} FAIL, {sum(r['pass'] is None for r in rows)} N/A"
    )
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main())
