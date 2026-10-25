import sys, zipfile, io
from pypdf import PdfReader
import openpyxl
def pdf(b):
    r=PdfReader(io.BytesIO(b)); t=(r.pages[0].extract_text() or '')[:160].replace('\n',' | ')
    return f"PDF pages={len(r.pages)} p1='{t}'"
def xlsx(b):
    wb=openpyxl.load_workbook(io.BytesIO(b)); return "XLSX sheets="+",".join(f"{ws.title}({ws.max_row}x{ws.max_column})" for ws in wb.worksheets)
def md(b):
    return "MD first: "+" / ".join(b.decode('utf-8').splitlines()[:4])
def one(name,b):
    if b[:4]==b'%PDF': return pdf(b)
    if name.endswith('.xlsx'): return xlsx(b)
    if name.endswith('.md'): return md(b)
    return f"magic={b[:4]!r}"
for f in sys.argv[1:]:
    b=open(f,'rb').read()
    if f.endswith('.zip'):
        z=zipfile.ZipFile(f); print(f.split('/')[-1],'ZIP:')
        for n in z.namelist(): print('   ',n,len(z.read(n)),one(n,z.read(n)))
    else: print(f.split('/')[-1], one(f,b))
