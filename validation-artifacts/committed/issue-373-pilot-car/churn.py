# #373 Rule 5 churn: hash every surface for each body, on the root given.
# Usage: python churn.py <repo root>   (run once on the base, once on the branch)
import sys, os, io, copy, hashlib, json, re
root = sys.argv[1]; sys.path.insert(0, root); sys.path.insert(1, os.getcwd())
os.environ["RENDER_API_SECRET"] = "s"
from fastapi.testclient import TestClient
from src.api.render_api import app
from tests.test_plan_sheet_device_summary import FLAGGER_BODY, SHOULDER_BODY
from tests.test_r117_jurisdiction_raised import _broadway_denver
import src.api.render_api as r; assert os.path.normcase(r.__file__).startswith(os.path.normcase(os.path.normpath(root))), r.__file__
from openpyxl import load_workbook
import pypdfium2 as pdfium
c = TestClient(app); H = {"Authorization": "Bearer s"}
bodies = {}
for k in (None, "denver", "greeley"):
    for car in (False, True):
        b = copy.deepcopy(FLAGGER_BODY); b["pilotCar"] = car
        if k: b["jurisdiction_key"] = k
        if k == "greeley": b["street_class"] = "arterial"
        bodies[f"flagger-{k}-pilot{car}"] = b
    s = copy.deepcopy(SHOULDER_BODY)
    if k: s["jurisdiction_key"] = k
    if k == "greeley": s["street_class"] = "arterial"
    bodies[f"shoulder-{k}"] = s
bodies["ni-denver"] = _broadway_denver()
# F12 (the #370 row) needs a live site scan, so it is checked on production.
TS = re.compile(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?(\.\d+)?")
def xl(content):
    wb = load_workbook(io.BytesIO(content)); out = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb]; wb.close()
    return TS.sub("TS", json.dumps(out, default=str))
def post(path, body):
    resp = c.post(path, json=body, headers=H); assert resp.status_code == 200, (path, resp.text[:300]); return resp
def page1(content):
    pdf = pdfium.PdfDocument(content)
    try: return TS.sub("TS", pdf[0].get_textpage().get_text_range())
    finally: pdf.close()
out = {}
for name, b in bodies.items():
    row = {}
    row["page1"] = page1(post("/render/pdf", b).content)
    row["md"] = post("/render/markdown", b).text
    row["xlsx"] = xl(post("/render/xlsx", b).content)
    row["quote"] = xl(post("/render/quote", {"scenario": b}).content)
    row["quote_json"] = post("/render/quote-breakdown", {"scenario": b}).text
    row["breakdown"] = post("/render/device-breakdown", b).text
    row["audit"] = post("/render/audit", b).text
    row["audit+breakdown"] = post("/render/audit", {**b, "include_breakdown": True}).text
    out[name] = {k: hashlib.sha256(v.encode()).hexdigest()[:12] for k, v in row.items()}
print(json.dumps(out))
