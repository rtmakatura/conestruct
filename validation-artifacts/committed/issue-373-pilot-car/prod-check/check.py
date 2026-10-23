"""#373 production check: the pilot car is listed on every surface that lists equipment.

Run from the repo root with GATE_BYPASS_TOKEN set:
``uv run python validation-artifacts/committed/issue-373-pilot-car/prod-check/check.py``.
It replays the audit lane's own requests (#370, ``audit-lane/runs/<id>/request.json``)
through www.conestruct.com's ``/api/render/*`` proxy: the audit (with its
breakdown) for a scan token, then page 1, the device XLSX, the quote XLSX, the
quote breakdown, the device breakdown and the crew sheet carrying it.  The
token stays in memory and is never written (#371).  Writes ``results.json``.

Rows: F12 (flagger, ``pilotCar: true``, the #373 acceptance) and F1 (flagger,
no pilot car, the control: no pilot-car row anywhere).
"""

from __future__ import annotations

import io
import json
import re
import sys
import time
from datetime import UTC, datetime
from pathlib import Path

import httpx
import pypdfium2 as pdfium
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
sys.path.insert(0, str(REPO / "scripts"))
from gate import gate_headers  # noqa: E402

BASE = "https://www.conestruct.com"
RUNS = REPO / "validation-artifacts/committed/audit-lane/runs"
NAME = 'Pilot car with G20-4 "PILOT CAR/FOLLOW ME" sign'
GAP_S = 4.0


class Paced:
    def __init__(self) -> None:
        self.c = httpx.Client(
            headers={**gate_headers(BASE), "content-type": "application/json"}, timeout=300
        )
        self.last = 0.0

    def post(self, path: str, body: dict) -> httpx.Response:
        for attempt in range(4):
            wait = self.last + GAP_S - time.time()
            if wait > 0:
                time.sleep(wait)
            r = self.c.post(f"{BASE}{path}", content=json.dumps(body))
            self.last = time.time()
            print(time.strftime("%H:%M:%S"), path, r.status_code, flush=True)
            if r.status_code == 429:
                time.sleep(65)
                continue
            if r.status_code == 400 and "site_scan_unavailable" in r.text[:2000]:
                time.sleep(30 * 2**attempt)
                continue
            return r
        return r


def page1(content: bytes) -> dict:
    pdf = pdfium.PdfDocument(content)
    try:
        lines = pdf[0].get_textpage().get_text_range().splitlines()
    finally:
        pdf.close()
    s = lines.index("TRAFFIC CONTROL DEVICE SUMMARY") + 2
    e = next(i for i in range(s, len(lines)) if lines[i].startswith("TOTAL DEVICES"))
    rows = [ln.removeprefix("— ") for ln in lines[s:e]]
    return {
        "pilot_rows": [r for r in rows if r.startswith("Pilot car")],
        "total": int(lines[e].removeprefix("TOTAL DEVICES ")),
    }


def device_xlsx(content: bytes) -> dict:
    wb = load_workbook(io.BytesIO(content))
    try:
        rows = list(wb["Device List"].iter_rows(min_row=2, values_only=True))
        total = dict(wb["Summary"].iter_rows(min_row=2, values_only=True))["Total device count"]
    finally:
        wb.close()
    return {
        "pilot_rows": [list(r[1:]) for r in rows if r[1] == NAME],
        "any_pilot_text": any("Pilot car" in str(c) for r in rows for c in r),
        "total": int(total),
    }


def quote_xlsx(content: bytes) -> dict:
    wb = load_workbook(io.BytesIO(content))
    try:
        rows = list(wb["Equipment Detail"].iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    return {"pilot_rows": [list(r[1:]) for r in rows if r[1] == NAME]}


def check_row(p: Paced, rid: str) -> dict:
    sc = json.loads((RUNS / rid / "request.json").read_text("utf-8"))["scenario"]
    out: dict = {"kind": sc["kind"], "pilotCar": sc.get("pilotCar")}
    a = p.post(
        "/api/render/audit",
        {"scenario": {**sc, "include_scan_token": True, "include_breakdown": True}},
    )
    out["audit_status"] = a.status_code
    if a.status_code != 200:
        out["audit_error"] = a.text[:500]
        return out
    audit = a.json()
    token = audit.get("site_scan_token")
    out["audit_breakdown_pilot"] = [
        d for d in (audit.get("breakdown") or {}).get("devices", []) if d["device"] == NAME
    ]
    out["audit_pilot_car_text"] = (audit.get("sections", {}).get("flagger") or {}).get(
        "pilot_car_text"
    )
    dl = {**sc}
    if token:
        dl["site_scan"] = {"proceed_if_unavailable": False, "token": token}
    body = {"scenario": dl}
    for key, path, parse in (
        ("page1", "/api/render/pdf", page1),
        ("device_xlsx", "/api/render/xlsx", device_xlsx),
        ("quote_xlsx", "/api/render/quote", quote_xlsx),
    ):
        r = p.post(path, body)
        out[f"{key}_status"] = r.status_code
        if r.status_code == 200:
            out[key] = parse(r.content)
    r = p.post("/api/render/quote-breakdown", body)
    out["quote_breakdown_status"] = r.status_code
    if r.status_code == 200:
        q = r.json()
        out["quote_breakdown"] = {
            "pilot_lines": [ln for ln in q["equipment_lines"] if ln["device_type"] == NAME],
            "equipment_total": q["equipment_total"],
            "total": q["total"],
        }
    r = p.post("/api/render/device-breakdown", body)
    out["breakdown_status"] = r.status_code
    if r.status_code == 200:
        b = r.json()
        out["breakdown"] = {
            "pilot_rows": [d for d in b["devices"] if d["device"] == NAME],
            "total_devices": b["total_devices"],
        }
    r = p.post("/api/render/markdown", body)
    out["crew_status"] = r.status_code
    if r.status_code == 200:
        out["crew_pilot_lines"] = [ln for ln in r.text.splitlines() if re.search(r"Pilot car", ln)]
    return out


def verdict(rid: str, r: dict) -> list[str]:
    keys = [k for k in r if k.endswith("_status")]
    if any(r[k] != 200 for k in keys):
        return [f"status {[(k, r[k]) for k in keys if r[k] != 200]}"]
    bad = []
    want = 1 if r["pilotCar"] else 0
    counts = {
        "page1": len(r["page1"]["pilot_rows"]),
        "device_xlsx": len(r["device_xlsx"]["pilot_rows"]),
        "quote_xlsx": len(r["quote_xlsx"]["pilot_rows"]),
        "quote_breakdown": len(r["quote_breakdown"]["pilot_lines"]),
        "breakdown": len(r["breakdown"]["pilot_rows"]),
        "audit_breakdown": len(r["audit_breakdown_pilot"]),
        "crew": len(r["crew_pilot_lines"]),
    }
    bad += [f"{k}: {n} pilot-car rows, want {want}" for k, n in counts.items() if n != want]
    if r["page1"]["total"] != r["device_xlsx"]["total"]:
        bad.append(f"page 1 total {r['page1']['total']} != XLSX {r['device_xlsx']['total']}")
    if r["breakdown"]["total_devices"] != r["device_xlsx"]["total"]:
        bad.append("breakdown total != XLSX total")
    if want:
        x = r["device_xlsx"]["pilot_rows"][0]
        if x[2:5] != ["—", "EACH", 1] or "S-630-1 Sheet 26" not in (x[5] or ""):
            bad.append(f"device XLSX row {x}")
        q = r["quote_xlsx"]["pilot_rows"][0]
        if (q[3], q[5], q[7]) != (1, "—", "—"):
            bad.append(f"quote XLSX row {q}")
        if r["quote_breakdown"]["pilot_lines"][0].get("unpriced") is not True:
            bad.append("quote breakdown line not unpriced")
        if r["audit_breakdown_pilot"] != r["breakdown"]["pilot_rows"]:
            bad.append("audit breakdown != device breakdown")
        if r["audit_pilot_car_text"] != "1 vehicle":
            bad.append(f"audit pilot_car_text {r['audit_pilot_car_text']!r}")
    elif r["device_xlsx"]["any_pilot_text"]:
        bad.append("control XLSX mentions a pilot car")
    return bad


def main() -> None:
    p = Paced()
    healthz = httpx.get(
        "https://rtmakatura--conestruct-render-fastapi-app.modal.run/healthz", timeout=60
    ).json()
    results: dict = {
        "base": BASE,
        "checked_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "healthz_sha": healthz.get("sha") or healthz,
        "rows": {},
    }
    for rid in ("F12", "F1"):
        r = check_row(p, rid)
        r["fail"] = verdict(rid, r)
        r["verdict"] = "FAIL" if r["fail"] else "PASS"
        results["rows"][rid] = r
        print(rid, r["verdict"], r["fail"], flush=True)
    passed = sum(r["verdict"] == "PASS" for r in results["rows"].values())
    results["summary"] = f"{passed}/{len(results['rows'])} PASS"
    (HERE / "results.json").write_text(
        json.dumps(results, indent=1, ensure_ascii=False) + "\n", "utf-8"
    )
    print(results["summary"])


if __name__ == "__main__":
    main()
