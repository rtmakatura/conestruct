"""s2-arc24 live check, API leg: the deployed proxies at www.conestruct.com
(backend Modal 196c116) — the deliverables at the Denver pin with and
without a jurisdiction key."""

import io
import re
import sys

import httpx
import pypdfium2 as pdfium
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
BASE = "https://www.conestruct.com"
print("healthz:", httpx.get("https://rtmakatura--conestruct-render-fastapi-app.modal.run/healthz", timeout=30).json())
DENVER = {
    "kind": "shoulder",
    "meta": {"project": "", "address": "", "lat": 39.7269, "lng": -104.9873, "bearingDeg": 180},
    "roadType": "rural_divided", "speed": 65, "lanes": 2, "laneWidth": 12, "divided": True,
    "workType": "utility_locate", "duration": "short", "workLen": 1000, "night": False,
    "jurisdiction_key": "denver",
}
NO_KEY = {k: v for k, v in DENVER.items() if k != "jurisdiction_key"}


def pdf_pages(b):
    d = pdfium.PdfDocument(b)
    out = []
    for p in d:
        tp = p.get_textpage()
        out.append(tp.get_text_range(0, tp.count_chars()))
    d.close()
    return out


for label, sc in (("DENVER", DENVER), ("NO KEY", NO_KEY)):
    print("==", label)
    r = httpx.post(f"{BASE}/api/render/xlsx", json={"scenario": sc}, timeout=120)
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    s = {str(a): b for a, b in wb["Summary"].iter_rows(values_only=True)}
    print("  XLSX Summary:", r.status_code, "Jurisdiction =", s["Jurisdiction"], "| Work zone length (ft) =", s["Work zone length (ft)"])
    md = httpx.post(f"{BASE}/api/render/markdown", json={"scenario": sc}, timeout=120).text
    print("  crew MD:", re.search(r"- \*\*Jurisdiction:\*\* .*", md).group(0), "|", re.search(r"- \*\*Work zone length:\*\* .*", md).group(0), "| step 3 station -", re.search(r"between station -([\d,]+) ft", md).group(1))
    crew = pdf_pages(httpx.post(f"{BASE}/api/render/crew-pdf", json={"scenario": sc}, timeout=120).content)
    print("  crew PDF:", re.search(r"Jurisdiction: [^\n]+", crew[0]).group(0).strip())
    plan = pdf_pages(httpx.post(f"{BASE}/api/render/pdf", json={"scenario": sc}, timeout=120).content)
    i = plan[1].find("CORRIDOR DETAILS")
    print("  plan p.2:", plan[1][i:i + 230].replace("\r\n", " | "))
    print("  plan p.1 callout:", re.search(r"WORK ZONE = [^\n]+", plan[0]).group(0))
    a = httpx.post(f"{BASE}/api/render/audit", json={"scenario": sc}, timeout=120).json()
    print("  audit corridor_spec:", a["sections"]["corridor_spec"])
    cs = httpx.post(f"{BASE}/api/render/corridor-spec", json={"kind": "shoulder", "speed": 65, "roadType": "rural_divided"}, timeout=60).json()
    print("  corridor-spec (picker):", cs)
    q = httpx.post(f"{BASE}/api/render/quote", json={"scenario": sc}, timeout=120)
    qwb = load_workbook(io.BytesIO(q.content), read_only=True)
    rows = [r[0] for r in qwb["Quote Summary"].iter_rows(min_row=6, max_row=7, max_col=1, values_only=True)]
    print("  quote header:", rows)
