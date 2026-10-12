"""#314 (R136): everything but page 1 stays byte-identical.

Prints a digest per deliverable for the four shoulder plans with and without
a sidewalk, from the given checkout: the placements (every device's type,
station, offset, label), the audit JSON, the crew markdown, the device
breakdown (bytes); the XLSX cells (minus the generated stamp); the audit PDF
and crew PDF page text.  It also prints page 1's device positions for two
non-shoulder plans (flagger, near-intersection), which no channelizer of
theirs should move.  Run on main and on the branch; diff the outputs.  The
main checkout loads a .env at import, so the Mapbox token is cleared after.
Usage: python identity_probe.py <repo-root>
"""

from __future__ import annotations

import copy
import datetime
import hashlib
import io
import json
import os
import sys
from pathlib import Path
from typing import Any

ROOT = Path(sys.argv[1]).resolve()
sys.path.insert(0, str(ROOT))
os.environ["RENDER_API_SECRET"] = "probe"
AUTH = {"Authorization": "Bearer probe"}

import pypdfium2 as pdfium  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pydantic import TypeAdapter  # noqa: E402

import src.rendering.plan_sheet as ps  # noqa: E402
from src.api import render_api  # noqa: E402
from src.api.schemas import Scenario  # noqa: E402
from src.rules import site_detection as sd  # noqa: E402

sd._overpass_request_with_fallback = lambda *_a, **_kw: ({"elements": []}, None)  # type: ignore[assignment]
os.environ.pop("MAPBOX_TOKEN", None)
client = TestClient(render_api.app)

rec: dict[str, Any] = {}
_orig_layout = ps._layout_device_positions


def _layout(placements, x_of, station_max_visible, params, shoulder_width_ft, site_flags=None):
    items, lights = _orig_layout(placements, x_of, station_max_visible, params, shoulder_width_ft, site_flags)
    rec["items"] = [(p.device_type.name, p.label, round(x, 3), round(y, 3)) for p, x, y in items]
    return items, lights


ps._layout_device_positions = _layout

FX = ROOT / "tests" / "fixtures"
bw = json.loads((FX / "corridor" / "broadway-sb.json").read_text(encoding="utf-8"))
m = bw["meta"]
cand = m["confirmedRoad"]["candidate"]
m["centerline"] = cand["geometry"]
m["roadDirection"] = {"osmBearingDeg": cand["bearing"] % 360, "oneway": cand["tags"]["oneway"]}
bw["carriageway"] = {"oneway": "yes", "highwayClass": "primary", "twinDistanceM": None, "twinSearched": True}
bw["divided"] = False
bw_left = copy.deepcopy(bw)
bw_left["meta"]["work"] = {"side": "left", "travel": "with_geometry"}


def worst(name: str) -> dict[str, Any]:
    return json.loads((FX / "pdf_worst_case" / name).read_text(encoding="utf-8"))["scenario"]


SHOULDER = {"broadway-right": bw, "broadway-left": bw_left,
            "two-way": worst("control-typical.json"), "divided": worst("adv-shoulder.json")}
OTHER = {"flagger": worst("adv-flagger.json"), "near-intersection": worst("adv-near-intersection.json")}


def h(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()[:16]


def post(path: str, body: dict[str, Any]) -> Any:
    return client.post(path, headers=AUTH, json=body)


def pdf_text(content: bytes) -> bytes:
    doc = pdfium.PdfDocument(content)
    return "\f".join(doc[i].get_textpage().get_text_range() for i in range(len(doc))).encode()


def xlsx_cells(content: bytes) -> bytes:
    wb = load_workbook(io.BytesIO(content))
    return repr([(ws.title, c.coordinate, str(c.value)) for ws in wb for row in ws.iter_rows()
                 for c in row if c.value is not None and not isinstance(c.value, datetime.datetime)]).encode()


for name, base in SHOULDER.items():
    for ped in (False, True):
        body = copy.deepcopy(base)
        body["meta"]["siteConditions"] = {**(body["meta"].get("siteConditions") or {}), "pedestrian_facility": ped}
        tag = f"{name}{' +sidewalk' if ped else ''}"
        placements, *_ = render_api._placements_for(TypeAdapter(Scenario).validate_python(copy.deepcopy(body)))
        pl = repr([(p.device_type.name, p.station_ft, p.offset_ft, p.label) for p in placements]).encode()
        rows = {
            "placements": h(pl),
            "audit": h(post("/render/audit", body).content),
            "markdown": h(post("/render/markdown", body).content),
            "device-breakdown": h(post("/render/device-breakdown", body).content),
            "xlsx (cells)": h(xlsx_cells(post("/render/xlsx", body).content)),
            "audit-pdf (text)": h(pdf_text(post("/render/audit-pdf", body).content)),
            "crew-pdf (text)": h(pdf_text(post("/render/crew-pdf", body).content)),
        }
        for k, v in rows.items():
            print(f"{tag:26} {k:18} {v}")

for name, body in OTHER.items():
    rec.clear()
    r = post("/render/pdf", body)
    print(f"{name:26} {'page-1 positions':18} {h(repr(rec.get('items')).encode()) if r.status_code == 200 else r.status_code}")
