"""#372 evidence: page 1's device summary vs the device XLSX, and churn.

Usage: uv run python check.py <tree root> <out.json>
Renders, offline through the real endpoints of the tree at <tree root>:
the 14 #370 near-intersection requests (audit-lane/runs/N*/request.json)
plus the flagger/shoulder test bodies.  For each it records page 1's
summary rows + total, the XLSX rows + total, and a sha of page 1's text
(the churn key: run once on the base tree, once on the fix).
"""
import copy, hashlib, io, json, os, sys
from pathlib import Path

out_path = Path(sys.argv[2]).resolve()
root = Path(sys.argv[1]).resolve(); sys.path.insert(0, str(root)); os.chdir(root)
os.environ["RENDER_API_SECRET"] = "s"
import pypdfium2 as pdfium
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import src.api.render_api as r
assert Path(r.__file__).resolve().is_relative_to(root), r.__file__
from tests.test_plan_sheet_device_summary import FLAGGER_BODY, SHOULDER_BODY

here = Path(__file__).resolve().parent
runs = here.parent / "audit-lane" / "runs"
c = TestClient(r.app); H = {"Authorization": "Bearer s"}
bodies = {d.name: json.loads((d / "request.json").read_text())["scenario"]
          for d in sorted(runs.glob("N*"), key=lambda p: int(p.name[1:]))}
for k, cls in [(None, None), ("denver", None), ("greeley", "arterial"), ("thornton", None)]:
    for name, base in (("flagger", FLAGGER_BODY), ("shoulder", SHOULDER_BODY)):
        b = copy.deepcopy(base)
        if k: b["jurisdiction_key"] = k
        if cls: b["street_class"] = cls
        bodies[f"{name}-{k}"] = b

def page1(body):
    resp = c.post("/render/pdf", json=body, headers=H)
    if resp.status_code != 200: return {"status": resp.status_code, "detail": resp.text[:200]}
    pdf = pdfium.PdfDocument(resp.content)
    try: text = pdf[0].get_textpage().get_text_range()
    finally: pdf.close()
    lines = text.splitlines()
    s = lines.index("TRAFFIC CONTROL DEVICE SUMMARY") + 2
    e = next(i for i in range(s, len(lines)) if lines[i].startswith("TOTAL DEVICES"))
    return {"rows": [l.rsplit(" ", 1)[0].removeprefix("— ") + " | " + l.rsplit(" ", 1)[1] for l in lines[s:e]],
            "total": int(lines[e].removeprefix("TOTAL DEVICES ")),
            "page1_text_sha": hashlib.sha256(text.encode()).hexdigest()[:12],
            "page1_text_sha_outside_summary": hashlib.sha256("\n".join(lines[:s] + lines[e + 1:]).encode()).hexdigest()[:12]}

def xlsx(body):
    resp = c.post("/render/xlsx", json=body, headers=H)
    if resp.status_code != 200: return {"status": resp.status_code}
    wb = load_workbook(io.BytesIO(resp.content))
    try:
        rows = [f"{x[2]} | {x[5]}" for x in wb["Device List"].iter_rows(min_row=2, values_only=True)]
        total = dict(wb["Summary"].iter_rows(min_row=2, values_only=True))["Total device count"]
    finally: wb.close()
    return {"rows": rows, "total": total}

out = {}
for name, b in bodies.items():
    p, x = page1(b), xlsx(b)
    agree = ("rows" in p and "rows" in x and p["total"] == x["total"]
             and [q.rsplit(" | ", 1)[1] for q in p["rows"]] == [q.rsplit(" | ", 1)[1] for q in x["rows"]])
    out[name] = {"page1": p, "xlsx": x, "agree": agree}
    print(name, "agree" if agree else "DISAGREE", p.get("total"), x.get("total"), flush=True)
out_path.write_text(json.dumps(out, indent=1, ensure_ascii=False), encoding="utf-8")
