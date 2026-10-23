"""#372 production check: page 1's device summary equals the device XLSX.

Run from the repo root with GATE_BYPASS_TOKEN set:
``uv run python validation-artifacts/committed/issue-372-ni-summary/prod-check/check.py``.
It replays the audit lane's own requests (#370, ``audit-lane/runs/<id>/request.json``)
through www.conestruct.com's ``/api/render/*`` proxy: the audit for a scan
token, then the plan PDF and device XLSX carrying it.  The token stays in
memory and is never written (#371).  Writes ``results.json`` beside itself.

Rows: N1-N14 (#372 acceptance) and F1 (flagger control, which already agreed).
PASS = page 1's summary quantities (in order) and total equal the XLSX's
Device List quantities and "Total device count".
"""

from __future__ import annotations

import io
import json
import sys
import time
from datetime import UTC, datetime
from pathlib import Path

import httpx
import pypdfium2 as pdfium
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
sys.path.insert(0, str(REPO / "scripts"))
from gate import gate_headers  # noqa: E402

BASE = "https://www.conestruct.com"
RUNS = REPO / "validation-artifacts/committed/audit-lane/runs"
ROWS = [f"N{i}" for i in range(1, 15)] + ["F1"]
GAP_S = 4.0


class Paced:
    def __init__(self) -> None:
        self.c = httpx.Client(
            headers={**gate_headers(BASE), "content-type": "application/json"}, timeout=300
        )
        self.last = 0.0

    def post(self, path: str, body: dict) -> httpx.Response:
        for attempt in range(4):
            wait = self.last + GAP_S - time.time()
            if wait > 0:
                time.sleep(wait)
            r = self.c.post(f"{BASE}{path}", content=json.dumps(body))
            self.last = time.time()
            print(time.strftime("%H:%M:%S"), path, r.status_code, flush=True)
            if r.status_code == 429:
                time.sleep(65)
                continue
            if r.status_code == 400 and "site_scan_unavailable" in r.text[:2000]:
                time.sleep(30 * 2**attempt)
                continue
            return r
        return r


def summary(content: bytes) -> dict:
    pdf = pdfium.PdfDocument(content)
    try:
        lines = pdf[0].get_textpage().get_text_range().splitlines()
    finally:
        pdf.close()
    s = lines.index("TRAFFIC CONTROL DEVICE SUMMARY") + 2
    e = next(i for i in range(s, len(lines)) if lines[i].startswith("TOTAL DEVICES"))
    return {
        "rows": [ln.removeprefix("— ") for ln in lines[s:e]],
        "total": int(lines[e].removeprefix("TOTAL DEVICES ")),
    }


def xlsx(content: bytes) -> dict:
    wb = load_workbook(io.BytesIO(content))
    try:
        rows = [f"{r[2]} {r[5]}" for r in wb["Device List"].iter_rows(min_row=2, values_only=True)]
        total = dict(wb["Summary"].iter_rows(min_row=2, values_only=True))["Total device count"]
    finally:
        wb.close()
    return {"rows": rows, "total": int(total)}


def check_row(p: Paced, rid: str) -> dict:
    sc = json.loads((RUNS / rid / "request.json").read_text("utf-8"))["scenario"]
    out: dict = {"kind": sc["kind"]}
    a = p.post("/api/render/audit", {"scenario": {**sc, "include_scan_token": True}})
    out["audit_status"] = a.status_code
    if a.status_code != 200:
        out["audit_error"] = a.text[:500]
        return out
    token = a.json().get("site_scan_token")
    dl = {**sc}
    if token:
        dl["site_scan"] = {"proceed_if_unavailable": False, "token": token}
    for kind, parse in (("pdf", summary), ("xlsx", xlsx)):
        r = p.post(f"/api/render/{kind}", {"scenario": dl})
        out[f"{kind}_status"] = r.status_code
        if r.status_code == 200:
            out[kind] = parse(r.content)
    return out


def verdict(r: dict) -> list[str]:
    if any(r.get(k) != 200 for k in ("audit_status", "pdf_status", "xlsx_status")):
        return [f"status {[r.get(k) for k in ('audit_status', 'pdf_status', 'xlsx_status')]}"]
    bad = []
    if r["pdf"]["total"] != r["xlsx"]["total"]:
        bad.append(f"total {r['pdf']['total']} != {r['xlsx']['total']}")
    qty = lambda rows: [x.rsplit(" ", 1)[1] for x in rows]  # noqa: E731
    if qty(r["pdf"]["rows"]) != qty(r["xlsx"]["rows"]):
        bad.append("row quantities differ")
    return bad


def main() -> None:
    p = Paced()
    healthz = httpx.get(
        "https://rtmakatura--conestruct-render-fastapi-app.modal.run/healthz", timeout=60
    ).json()
    results: dict = {
        "base": BASE,
        "checked_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "healthz_sha": healthz.get("sha") or healthz,
        "rows": {},
    }
    for rid in ROWS:
        r = check_row(p, rid)
        r["fail"] = verdict(r)
        r["verdict"] = "FAIL" if r["fail"] else "PASS"
        results["rows"][rid] = r
        print(rid, r["verdict"], r.get("pdf", {}).get("total"), r.get("xlsx", {}).get("total"),
              r["fail"], flush=True)
    passed = sum(r["verdict"] == "PASS" for r in results["rows"].values())
    results["summary"] = f"{passed}/{len(results['rows'])} PASS"
    (HERE / "results.json").write_text(json.dumps(results, indent=1, ensure_ascii=False) + "\n", "utf-8")
    print(results["summary"])


if __name__ == "__main__":
    main()
