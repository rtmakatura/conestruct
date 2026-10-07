"""R108 / R110 -- the Rule 5 outcomes, measured against checkpoint.md section 4.

Renders one E Colfax shoulder plan (the arc's standing spot, inside Denver)
three ways through the real API (FastAPI TestClient, no network: no work
side, so no corridor check and no scan):

  A  no jurisdiction, no street class        -- today's default
  B  Denver + Arterial, operator-set          -- a confirmed plan before R108
  C  Denver + Arterial, both R108 GUESSES     -- what a fresh pin now sends

and prints what each surface says.  Prediction 1 (checkpoint section 4)
is A vs C: a guessed jurisdiction turns its evaluation on.  B vs C isolates
what the guess record itself adds: input_guesses in the audit JSON and
audit PDF, the R110 Q3 clause on the XLSX and crew line, and nothing else.

Usage (repo root):  uv run --frozen --extra dev python
  validation-artifacts/committed/setup-what-redesign/probes/r108_churn.py
"""

from __future__ import annotations

import io
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("RENDER_API_SECRET", "probe-secret")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from src.api.render_api import app  # noqa: E402
from src.rendering.audit_blocks import audit_to_blocks  # noqa: E402

PIN = (39.7402, -104.956)  # E Colfax Ave, the declutter arc's spot
AUTH = {"Authorization": f"Bearer {os.environ['RENDER_API_SECRET']}"}


def body(**over):
    b = {
        "kind": "shoulder",
        "meta": {"project": "R108 probe", "address": "E Colfax Ave", "lat": PIN[0], "lng": PIN[1]},
        "roadType": "urban_arterial",
        "speed": 30,
        "lanes": 2,
        "laneWidth": 12.0,
        "divided": False,
        "workType": "utility_locate",
        "duration": "short",
        "workLen": 1000.0,
        "night": False,
    }
    b.update(over)
    return b


CASES = {
    "A_none": body(),
    "B_operator": body(jurisdiction_key="denver", street_class="arterial"),
    "C_guessed": body(
        jurisdiction_key="denver",
        street_class="arterial",
        guesses={
            "street_class": {"highwayClass": "primary"},
            "jurisdiction_key": {"lat": PIN[0], "lng": PIN[1]},
        },
    ),
}


def main() -> None:
    client = TestClient(app)
    out = {}
    for name, b in CASES.items():
        bd = client.post("/render/device-breakdown", json=b, headers=AUTH)
        au = client.post("/render/audit", json=b, headers=AUTH)
        xl = client.post("/render/xlsx", json=b, headers=AUTH)
        md = client.post("/render/markdown", json=b, headers=AUTH)
        assert bd.status_code == au.status_code == xl.status_code == md.status_code == 200, (
            name,
            bd.text[:200],
            au.text[:200],
        )
        bdj, auj = bd.json(), au.json()
        jur = bdj.get("jurisdiction") or {}
        summary = {
            r[0].value: r[1].value
            for r in load_workbook(io.BytesIO(xl.content))["Summary"].iter_rows()
            if r[0].value
        }
        crew_line = next(
            (ln for ln in md.text.splitlines() if ln.startswith("- **Jurisdiction:**")), None
        )
        blocks = repr(audit_to_blocks(auj))
        out[name] = {
            "total_devices": bdj.get("total_devices"),
            "unique_types": bdj.get("unique_types"),
            "jurisdiction_block": bool(jur),
            "applied_deltas": len(jur.get("applied_deltas", []) or []),
            "hours_eval_status": (jur.get("hours_eval") or {}).get("status"),
            "permit_tier": ((jur.get("permit") or {}).get("tier_suggested")),
            "pending_verification_count": auj["pending_verification"]["count"],
            "plan_flags": auj["plan_flags"],
            "input_guesses": auj.get("input_guesses"),
            "xlsx_jurisdiction": summary.get("Jurisdiction"),
            "crew_jurisdiction_line": crew_line,
            "audit_pdf_guess_block": "Guessed inputs" in blocks,
        }
    print(json.dumps(out, indent=2, default=str))

    a, b, c = out["A_none"], out["B_operator"], out["C_guessed"]
    print("\nB vs C -- what the guess RECORD adds (everything else must be equal):")
    for k in a:
        if b[k] != c[k]:
            print(f"  {k}: B={b[k]!r}  C={c[k]!r}")
    print("\nA vs C -- what a guessed jurisdiction + class turns on (prediction 1):")
    for k in a:
        if a[k] != c[k]:
            print(f"  {k}: A={a[k]!r}  C={c[k]!r}")


if __name__ == "__main__":
    main()
