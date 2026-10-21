"""#363 byte-identity: hash every surface over a reduced-speed shoulder grid.

Run from the worktree root: PYTHONPATH=. uv run --extra dev python <this> <out.json>
Audit JSON (canonical), audit PDF text, plan PDF text, crew markdown and
XLSX cell values per plan (timestamps masked).  Compare a base run with an after run.
"""
import datetime, hashlib, io, itertools, json, sys
from openpyxl import load_workbook
from tests.s630._harness import _common
from tests.s630 import test_cross_surface as tcs

import os
from fastapi.testclient import TestClient
os.environ["RENDER_API_SECRET"] = "test-secret-do-not-deploy"
os.environ["MAPBOX_TOKEN"] = "test-token"
tcs._ps.httpx.get = lambda *a, **k: tcs._FakeTile()  # the fixture's offline tile
tcs._ps._validate_corridor_bearing = lambda corridor: None
api = TestClient(tcs._render_app)
H = tcs._API_HEADERS
h = lambda b: hashlib.sha256(b).hexdigest()[:16]
out = {}
for road, (s, wz), wl in itertools.product(
        ["urban_arterial", "rural_undivided", "rural_divided", "freeway"],
        [(45, 35), (55, 45), (55, 50), (65, 55), (65, 60), (75, 65), (75, 70), (55, None)],
        [200, 1000, 5000]):
    body = {**_common(s, wz), "roadType": road, "divided": road in ("rural_divided", "freeway"),
            "workLen": wl}
    key = f"{road} {s}->{wz} L{wl}"
    a = api.post("/render/audit", json=body, headers=H)
    if a.status_code != 200:
        out[key] = f"HTTP {a.status_code}"; continue
    aj = a.json()
    row = {"audit": h(json.dumps(aj, sort_keys=True).encode()),
           "note": any(i["kind"] == "fines_double_spacing"
                       for i in (aj.get("pending_verification") or {}).get("items", []))}
    ap = api.post("/render/audit-pdf", json=body, headers=H)
    row["audit_pdf"] = h("\n".join(tcs._pdf_text(ap.content)).encode()) if ap.status_code == 200 else ap.status_code
    pp = api.post("/render/pdf", json=body, headers=H)
    row["plan_pdf"] = h("\n".join(tcs._pdf_text(pp.content)).encode()) if pp.status_code == 200 else pp.status_code
    md = api.post("/render/markdown", json=body, headers=H)
    row["crew"] = h(md.content)
    xl = api.post("/render/xlsx", json=body, headers=H)
    if xl.status_code == 200:
        wb = load_workbook(io.BytesIO(xl.content))
        # The export timestamp (Summary!B8) is masked: it changes every run.
        cells = [[None if isinstance(c.value, datetime.datetime) else c.value for c in r]
                 for ws in wb for r in ws.iter_rows()]
        row["xlsx"] = h(repr(cells).encode())
    out[key] = row
json.dump(out, open(sys.argv[1], "w"), indent=1, sort_keys=True)
print(len(out), "plans;", sum(1 for v in out.values() if isinstance(v, dict) and v["note"]), "with the note")
