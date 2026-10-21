"""#365 byte-identity proof: dump every surface for five plans, offline.

Usage: python byte_identity.py <repo_root> <out.json>

Run once against the base tree and once against the branch, then diff the
two JSON files.  Each plan dumps the served audit, the corridor geometry the
picker draws, the crew markdown, the quote breakdown, the device breakdown,
the XLSX cell values, the plan sheet's page text and the Mapbox requests
page 2 makes (the drawn geometry).  Mapbox and Overpass go to stubs.
"""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path
from typing import Any

root = Path(sys.argv[1]).resolve()
sys.path.insert(0, str(root))

import os  # noqa: E402

os.environ["RENDER_API_SECRET"] = "test-secret"
os.environ["MAPBOX_TOKEN"] = "test-token"

import openpyxl  # noqa: E402
import pypdfium2 as pdfium  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from PIL import Image  # noqa: E402

import src  # noqa: E402
import src.api.render_api as render_api  # noqa: E402
from src.api import site_scan as ss  # noqa: E402
from src.rendering import plan_sheet as ps  # noqa: E402
from src.rules import site_detection as sd  # noqa: E402

assert Path(src.__file__).resolve().is_relative_to(root), src.__file__

AUTH = {"Authorization": "Bearer test-secret"}
PIN = {"lat": 39.73997, "lng": -104.96632, "bearingDeg": 90.0}
requests: list[str] = []


class _Png:
    def __init__(self) -> None:
        buf = io.BytesIO()
        Image.new("RGB", (4, 4), (128, 128, 128)).save(buf, format="PNG")
        self.content = buf.getvalue()

    def raise_for_status(self) -> None:
        return None


def _get(url: str, params: dict[str, str] | None = None, **_k: Any) -> _Png:
    requests.append(url)
    return _Png()


ps.httpx.get = _get  # type: ignore[assignment]
ps._validate_corridor_bearing = lambda corridor: None  # type: ignore[assignment]
sd._overpass_request_with_fallback = lambda q, **_k: (None, "stub: mirrors down")  # type: ignore[assignment]
render_api.ENABLED_SCENARIOS = frozenset({"shoulder", "flagger_lane_closure", "near_intersection"})

fixtures = root / "tests" / "fixtures" / "pdf_worst_case"


def _fixture(name: str) -> dict[str, Any]:
    s = json.loads((fixtures / f"{name}.json").read_text(encoding="utf-8"))["scenario"]
    return {**s, "meta": {**s["meta"], **PIN}}


def _ni(along: float) -> dict[str, Any]:
    from tests.test_lane_confidence_block import _ni_body

    b = _ni_body(alongStationFt=along)
    return {**b, "meta": {**b["meta"], **PIN}}


PLANS = {
    "ni_near_corner": _ni(-60.0),
    "ni_near_midblock": _ni(-600.0),
    "ni_far_side": _ni(800.0),
    "shoulder": _fixture("adv-shoulder"),
    "flagger": _fixture("adv-flagger"),
}


def _dump(client: TestClient, body: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    out["audit"] = client.post("/render/audit", json=body, headers=AUTH).json()
    out["corridor_geometry"] = client.post(
        "/render/corridor-geometry", json=body, headers=AUTH
    ).json()
    out["markdown"] = client.post("/render/markdown", json=body, headers=AUTH).text
    out["quote_breakdown"] = client.post(
        "/render/quote-breakdown", json={"scenario": body}, headers=AUTH
    ).json()
    out["device_breakdown"] = client.post(
        "/render/device-breakdown", json=body, headers=AUTH
    ).json()
    x = client.post("/render/xlsx", json=body, headers=AUTH)
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    out["xlsx"] = {
        ws.title: [
            [None if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)
        ]
        for ws in wb.worksheets
    }
    requests.clear()
    pdf = client.post("/render/pdf", json=body, headers=AUTH)
    doc = pdfium.PdfDocument(pdf.content)
    out["pdf_pages"] = [doc[i].get_textpage().get_text_range() for i in range(len(doc))]
    out["mapbox_requests"] = list(requests)
    return out


def main() -> None:
    ss.clear_memo()
    client = TestClient(render_api.app)
    result = {name: _dump(client, body) for name, body in PLANS.items()}
    Path(sys.argv[2]).write_text(json.dumps(result, indent=1, sort_keys=True), encoding="utf-8")


main()
