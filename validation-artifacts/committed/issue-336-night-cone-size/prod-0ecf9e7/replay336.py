"""#336 production check: replay audit-shoulder rows through production's
/api/render/* and read the cone label off every surface.

    python replay336.py

The gate token comes from GATE_BYPASS_TOKEN, else conestruct/site/.env.local;
it is never printed.  Same pacing and scan-token reuse as
audit-shoulder/fetch.py: one /api/render/audit per scenario (with the
breakdown and a signed scan token), then the five downloads on that token.
"""

from __future__ import annotations

import copy
import io
import json
import os
import sys
import time
from pathlib import Path

import httpx
import pypdfium2 as pdfium
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
RUNS = REPO / "validation-artifacts" / "committed" / "audit-shoulder" / "runs"
BASE = "https://www.conestruct.com"
GAP_S = 4.0
C36, C28 = "Traffic Cone (36-inch)", "Traffic Cone (28-inch)"

# (case id, audit row, overrides, expected label)
CASES = [
    ("bway-l-30-night", "bway-l-30-night", {}, C36),
    ("colo-r-40-night", "colo-r-40-night", {}, C36),
    ("expo-div-r-35-night", "expo-div-r-35-night", {}, C36),
    ("bway-l-30-day", "bway-l-30-night", {"night": False}, C28),
]


def _token() -> str:
    tok = os.environ.get("GATE_BYPASS_TOKEN")
    if tok:
        return tok
    for line in (REPO / "conestruct" / "site" / ".env.local").read_text("utf-8").splitlines():
        if line.startswith("GATE_BYPASS_TOKEN="):
            return line.split("=", 1)[1].strip().strip('"')
    raise SystemExit("GATE_BYPASS_TOKEN not found")


class Poster:
    def __init__(self) -> None:
        self.c = httpx.Client(
            timeout=180,
            headers={"x-conestruct-gate": _token(), "content-type": "application/json"},
        )
        self.last = 0.0

    def post(self, path: str, body: dict) -> httpx.Response:
        for attempt in range(4):
            wait = self.last + GAP_S - time.time()
            if wait > 0:
                time.sleep(wait)
            r = self.c.post(f"{BASE}{path}", content=json.dumps(body))
            self.last = time.time()
            print(f"  {path} -> {r.status_code}", flush=True)
            if r.status_code == 429:
                time.sleep(65)
                continue
            if r.status_code == 400 and "site_scan_unavailable" in r.text[:2000]:
                time.sleep(30 * 2**attempt)
                continue
            return r
        return r


def _pdf_pages(b: bytes) -> list[str]:
    doc = pdfium.PdfDocument(b)
    try:
        out = []
        for page in doc:
            tp = page.get_textpage()
            out.append(tp.get_text_range(0, tp.count_chars()))
        return out
    finally:
        doc.close()


def _rows(b: bytes, sheet: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(b), read_only=True)
    rows = list(wb[sheet].iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows


def _found(text: str) -> list[str]:
    return [s for s in (C36, C28) if s in text]


def main() -> int:
    p = Poster()
    results = {}
    for case, row, overrides, expected in CASES:
        print(case, flush=True)
        sc = {**json.loads((RUNS / row / "request.json").read_text("utf-8"))["scenario"], **overrides}
        audit_sc = {**copy.deepcopy(sc), "include_scan_token": True, "include_breakdown": True}
        a = p.post("/api/render/audit", {"scenario": audit_sc})
        surfaces: dict[str, list[str]] = {}
        if a.status_code == 200:
            aj = a.json()
            dl = copy.deepcopy(sc)
            if aj.get("site_scan_token"):
                dl["site_scan"] = {"proceed_if_unavailable": False, "token": aj["site_scan_token"]}
            devices = (aj.get("breakdown") or {}).get("devices") or []
            surfaces["breakdown"] = [str(d["device"]) for d in devices if "Cone" in str(d["device"])]
            got = {k: p.post(f"/api/render/{k}", {"scenario": dl}) for k in
                   ("pdf", "xlsx", "markdown", "crew-pdf", "quote")}
            bad = {k: r.status_code for k, r in got.items() if r.status_code != 200}
            if bad:
                surfaces["error"] = [json.dumps(bad)]
            else:
                pages = _pdf_pages(got["pdf"].content)
                (HERE / f"{case}-page1-text.txt").write_text(pages[0], "utf-8")
                (HERE / f"{case}-crew.md").write_text(got["markdown"].text, "utf-8")
                surfaces["page1"] = _found(pages[0])
                surfaces["crew_md"] = _found(got["markdown"].text)
                surfaces["crew_pdf"] = _found("\n".join(_pdf_pages(got["crew-pdf"].content)))
                surfaces["xlsx"] = [str(r[2]) for r in _rows(got["xlsx"].content, "Device List")
                                    if r[1] == "CONE"]
                surfaces["quote"] = [str(r[3]) for r in _rows(got["quote"].content, "Equipment Detail")
                                     if isinstance(r[0], int) and r[1] == "CONE"]
        else:
            surfaces["error"] = [f"audit {a.status_code}"]
        ok = "error" not in surfaces and all(v and set(v) == {expected} for v in surfaces.values())
        results[case] = {"expected": expected, "pass": ok, "surfaces": surfaces}
        print(f"  {'PASS' if ok else 'FAIL'} {surfaces}", flush=True)
    (HERE / "results.json").write_text(json.dumps(results, indent=1), "utf-8")
    return 0 if all(r["pass"] for r in results.values()) else 1


if __name__ == "__main__":
    sys.exit(main())
