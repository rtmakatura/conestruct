"""#333 evidence: page 1, the audit's spacing lines and crew step 3 for
three shoulder plans from the shoulder-work audit, rendered locally with
the network stubbed.  Writes <name>.png beside this file (or, with
--out DIR, into DIR) and prints the counts.

Run from a checkout's root:
    python -B validation-artifacts/committed/issue-333-buffer-devices/render333.py [--out DIR]
"""

from __future__ import annotations

import copy
import io
import json
import os
import re
import sys
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, os.getcwd())
os.environ.setdefault("RENDER_API_SECRET", "test-secret-do-not-deploy")

import pypdfium2 as pdfium  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from src.api import render_api  # noqa: E402
from src.rules import site_detection as sd  # noqa: E402

sd._overpass_request_with_fallback = lambda *a, **k: ({"elements": []}, None)  # type: ignore[assignment]
HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[sys.argv.index("--out") + 1]) if "--out" in sys.argv else HERE
OUT.mkdir(parents=True, exist_ok=True)
RUNS = Path("validation-artifacts/committed/audit-shoulder/runs")
AUTH = {"Authorization": f"Bearer {os.environ['RENDER_API_SECRET']}"}


def body(row: str) -> dict:
    s = json.loads((RUNS / row / "request.json").read_text(encoding="utf-8"))["scenario"]
    s = copy.deepcopy(s)
    s.pop("site_scan", None)
    cand = s["meta"]["confirmedRoad"]["candidate"]
    s["meta"]["centerline"] = cand["geometry"]
    s["meta"]["roadDirection"] = {
        "osmBearingDeg": cand["bearing"] % 360,
        "oneway": (cand.get("tags") or {}).get("oneway"),
    }
    return s


c = TestClient(render_api.app)
for row in ("bway-r-30", "ogden-r-25-50ft", "hartsel-r-65"):
    b = body(row)
    pdf = c.post("/render/pdf", headers=AUTH, json=b)
    pdf.raise_for_status()
    doc = pdfium.PdfDocument(pdf.content)
    doc[0].render(scale=150 / 72).to_pil().convert("RGB").save(OUT / f"{row}.png", optimize=True)
    text = doc[0].get_textpage().get_text_range()
    a = c.post("/render/audit", headers=AUTH, json=b).json()
    sp = a["sections"]["spacing"]
    md = c.post("/render/markdown", headers=AUTH, json=b).text
    xl = load_workbook(io.BytesIO(c.post("/render/xlsx", headers=AUTH, json=b).content))
    cones_xlsx = sum(
        r[5]
        for r in xl["Device List"].iter_rows(min_row=2, values_only=True)
        if r[2] and "Cone" in str(r[2])
    )
    label = re.search(r"BUFFER = [^\n]*", text)
    summary = re.search(r"Traffic Cone \(\d+-inch\) (\d+)", text)
    step = re.search(r"Then continue the line through the buffer space:[^\n]*", md)
    print(f"== {row}")
    print("  page 1 label:", label.group(0) if label else None)
    print(
        "  page 1 summary cones:",
        summary.group(1) if summary else None,
        "| xlsx cones:",
        cones_xlsx,
    )
    print("  audit:", sp.get("buffer_count_text"))
    print("  audit:", sp.get("downstream_count_text"))
    print("  crew:", step.group(0)[:220] if step else None)
