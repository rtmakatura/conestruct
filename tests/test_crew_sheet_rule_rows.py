"""#381: the crew sheet lists every row a jurisdiction rule adds to the
device XLSX, with the same quantity.

Thornton's stipulation 6 ("CONSTRUCTION AHEAD and MEN WORKING signs
required for each direction of travel affected", Thornton Construction &
TC Permit handout) adds "Advance Warning Signs" to every Thornton plan with
no placement.  Rule 11: the assertions read the real endpoints' output
(/render/markdown, /render/crew-pdf, /render/xlsx and the replication
snapshot) or the rendered crew sheet, not the row helper.
"""

from __future__ import annotations

import copy
import io
import re

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.render_api import app
from src.export.device_list import export_device_list
from src.narrative.crew_narrative import render_crew_narrative_markdown
from src.rules.device_aggregation import PILOT_CAR_NAME
from src.rules.devices import DeviceType
from tests.test_plan_sheet_device_summary import (
    FLAGGER_BODY,
    SHOULDER_BODY,
    _placements_for_body,
)
from tests.test_r117_jurisdiction_raised import _broadway_denver

AUTH = {"Authorization": "Bearer test-secret"}
AWS = "Advance Warning Signs (jurisdiction-required)"
RULE = (
    "CONSTRUCTION AHEAD and MEN WORKING signs required for each direction "
    "of travel affected (stipulation #6)."
)


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setenv("RENDER_API_SECRET", "test-secret")


@pytest.fixture()
def client():
    return TestClient(app)


def _with(base: dict, key: str | None = "thornton", **extra) -> dict:
    body = copy.deepcopy(base)
    if key:
        body["jurisdiction_key"] = key
    body.update(extra)
    return body


BODIES = {
    "shoulder": lambda: _with(SHOULDER_BODY),
    "flagger": lambda: _with(FLAGGER_BODY),
    "near_intersection": lambda: _with(_broadway_denver()),
}


def _post(client: TestClient, path: str, body: dict) -> bytes:
    resp = client.post(path, json=body, headers=AUTH)
    assert resp.status_code == 200, resp.text
    return resp.content


def _text(client: TestClient, path: str, body: dict) -> str:
    # The .md file is written in text mode: CRLF on Windows.
    return _post(client, path, body).decode().replace("\r\n", "\n")


def _rule_rows(content: bytes) -> list[tuple]:
    """Device XLSX rows a jurisdiction rule added with no placement."""
    wb = load_workbook(io.BytesIO(content))
    try:
        rows = wb["Device List"].iter_rows(min_row=2, values_only=True)
        return [r for r in rows if r and str(r[1] or "").endswith("(jurisdiction-required)")]
    finally:
        wb.close()


def _equipment(md: str) -> str:
    return md.split("## Required Equipment", 1)[1].split("## ", 1)[0]


@pytest.mark.parametrize("kind", sorted(BODIES))
def test_thornton_crew_sheet_lists_the_xlsx_rule_row(client, kind) -> None:
    body = BODIES[kind]()
    rows = _rule_rows(_post(client, "/render/xlsx", body))
    assert [r[1] for r in rows] == [AWS]
    md = _text(client, "/render/markdown", body)
    for r in rows:
        assert f"- {r[5]}× {r[1]}\n" in _equipment(md)


@pytest.mark.parametrize("kind", sorted(BODIES))
def test_thornton_setup_quotes_the_rule_and_leaves_the_position(client, kind) -> None:
    md = _text(client, "/render/markdown", BODIES[kind]())
    setup = md.split("## Setup Procedure", 1)[1].split("## Takedown", 1)[0]
    assert (
        f'   {AWS}. Thornton Construction & TC Permit handout, stipulation 6: "{RULE}"'
        " Not placed on the drawing: the rule gives no position. The Traffic"
        " Control Supervisor sets the position per the permit.\n"
    ) in setup
    takedown = md.split("## Takedown Procedure", 1)[1].split("## ", 1)[0]
    assert takedown.count(f"Also remove the {AWS}.") == 1


def test_thornton_crew_pdf_lists_the_rule_row(client, tmp_path) -> None:
    tmp = tmp_path / "crew.pdf"
    tmp.write_bytes(_post(client, "/render/crew-pdf", BODIES["shoulder"]()))
    pdf = pdfium.PdfDocument(str(tmp))
    try:
        text = " ".join(p.get_textpage().get_text_range() for p in pdf)
    finally:
        pdf.close()
    # #388: two signs for the one direction a shoulder plan affects.
    assert re.search(r"2× Advance Warning Signs\s+\(jurisdiction-required\)", text)


def test_a_topped_up_count_matches_the_xlsx(tmp_path) -> None:
    """A rule that raises a placed device's count (Thornton's night drum
    rule, were the layout short) sets the crew sheet's count too."""
    placements, params = _placements_for_body(_with(SHOULDER_BODY, key=None))
    assert any(p.device_type == DeviceType.DRUM for p in placements)
    deltas = [
        {
            "severity": "count",
            "status": "fires",
            "rule": "test: 500 drums",
            "effect": {"op": "add_device", "device": "drum", "qty": 500},
            "source": {"doc": "test record", "section": "1"},
        }
    ]
    path = tmp_path / "d.xlsx"
    export_device_list(placements, params, output_path=str(path), applied_deltas=deltas)
    wb = load_workbook(str(path))
    try:
        rows = [r for r in wb["Device List"].iter_rows(min_row=2, values_only=True) if r]
    finally:
        wb.close()
    drums = [r[5] for r in rows if r[1] == DeviceType.DRUM.value]
    assert sum(drums) == 500, drums
    md = render_crew_narrative_markdown(placements, params, applied_deltas=deltas)
    assert "- 500× Channelizing Drum (36-inch)\n" in _equipment(md)


def _section(snap: str, n: int) -> str:
    return snap.split(f"## {n} ", 1)[1].split(f"## {n + 1} ", 1)[0]


def test_snapshot_device_list_carries_the_rule_row(client) -> None:
    body = BODIES["shoulder"]()
    snap = _text(client, "/render/replication-snapshot", {"scenario": body})
    assert f"| {AWS} |" in _section(snap, 5)


@pytest.mark.parametrize(
    "body",
    [BODIES["shoulder"](), _with(FLAGGER_BODY, key="denver")],
    ids=["thornton-shoulder", "denver-flagger"],
)
def test_snapshot_crew_section_is_the_shipped_crew_sheet(client, body) -> None:
    snap = _text(client, "/render/replication-snapshot", {"scenario": body})
    md = _text(client, "/render/markdown", body)
    assert md in snap.split("## 6 Crew narrative", 1)[1]


def test_snapshot_device_list_carries_the_pilot_car(client) -> None:
    body = _with(FLAGGER_BODY, key=None, pilotCar=True)
    snap = _text(client, "/render/replication-snapshot", {"scenario": body})
    section5 = snap.split("## 5 Device list", 1)[1].split("## 6", 1)[0]
    assert f"| {PILOT_CAR_NAME} |" in section5


@pytest.mark.parametrize("kind", sorted(BODIES))
def test_no_rule_row_without_a_firing_rule(client, kind) -> None:
    # Shoulder and flagger with no jurisdiction; near-intersection on
    # Denver, whose arrow-board rule the layout's own board meets.
    body = _broadway_denver() if kind == "near_intersection" else BODIES[kind]()
    if kind != "near_intersection":
        body.pop("jurisdiction_key")
    md = _text(client, "/render/markdown", body)
    assert "jurisdiction-required" not in md
    assert "Not placed on the drawing" not in md
    assert "Also remove the" not in md
