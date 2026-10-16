"""#333: a shoulder plan channelizes its buffer at <= 2S.

CDOT S-630-1 Sheet 2 General Note 18b: "For tangents along the buffer space
or work area, spacing may not be greater than two times the speed limit"
(MUTCD 11th Ed. Sec 6K.01 para 04, p. 815, the same 2x).  Before #333 every
shoulder plan left the buffer bare: on N Broadway at 30 mph, cones ended at
the work zone (station 200) and the taper began at 400.
"""

from __future__ import annotations

import io
import os
import re
from collections.abc import Iterator

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api import render_api
from src.generation.layout import (
    buffer_cone_stations,
    generate_lane_closure_divided,
    generate_shoulder_closure_divided,
    generate_shoulder_closure_undivided,
)
from src.rules.devices import DeviceType
from src.rules.spacing import buffer_space
from src.rules.validators import ScenarioParams

_TEST_SECRET = "test-secret-do-not-deploy"
AUTH = {"Authorization": f"Bearer {_TEST_SECRET}"}


@pytest.fixture(scope="module", autouse=True)
def _render_secret() -> Iterator[None]:
    os.environ["RENDER_API_SECRET"] = _TEST_SECRET
    yield


def _params(speed: int, divided: bool, closure: str = "shoulder") -> ScenarioParams:
    return ScenarioParams(
        speed_mph=speed,
        num_lanes=2,
        closure_type=closure,
        road_type="rural",
        work_zone_length_ft=800.0,
        lane_width_ft=12.0,
        shoulder_width_ft=10.0 if divided else 8.0,
        is_divided=divided,
        jurisdiction="CDOT",
    )


def test_the_helper_takes_the_fewest_intervals_at_or_under_2s() -> None:
    # 200 ft at 30 mph: ceil(200 / 60) = 4 intervals of 50 ft -> 3 interior cones.
    assert buffer_cone_stations(200.0, 200.0, 30) == [250.0, 300.0, 350.0]
    # 645 ft at 65 mph: ceil(645 / 130) = 5 intervals of 129 ft -> 4 cones.
    assert buffer_cone_stations(1500.0, 645.0, 65) == pytest.approx(
        [1629.0, 1758.0, 1887.0, 2016.0]
    )
    # A buffer of 2S or less already meets Note 18b between its two end devices.
    assert buffer_cone_stations(100.0, 60.0, 30) == []
    assert buffer_cone_stations(100.0, 0.0, 30) == []


@pytest.mark.parametrize("speed", [25, 30, 35, 45, 55, 65, 75])
@pytest.mark.parametrize(
    "generator,divided",
    [(generate_shoulder_closure_divided, True), (generate_shoulder_closure_undivided, False)],
)
def test_no_gap_over_2s_from_the_work_zone_to_the_taper(generator, divided, speed) -> None:
    params = _params(speed, divided)
    placements = generator(params)
    wz = params.work_zone_length_ft
    buf = buffer_space(speed, jurisdiction="CDOT")
    # The lane-edge line from the work zone's downstream end to the taper's
    # downstream end: work-zone cones, buffer cones, then the last taper drum.
    line = sorted(
        p.station_ft
        for p in placements
        if p.device_type in (DeviceType.CONE, DeviceType.DRUM) and 0.0 <= p.station_ft <= wz + buf
    )
    assert line[0] == 0.0 and line[-1] == pytest.approx(wz + buf)
    gaps = [b - a for a, b in zip(line, line[1:], strict=False)]
    assert max(gaps) <= 2 * speed + 1e-6, gaps
    buffer_cones = [
        p for p in placements if p.device_type == DeviceType.CONE and wz < p.station_ft < wz + buf
    ]
    assert len(buffer_cones) == len(buffer_cone_stations(wz, buf, speed)) > 0


def test_a_lane_closure_is_not_changed() -> None:
    params = _params(55, True, closure="lane")
    placements = generate_lane_closure_divided(params)
    wz = params.work_zone_length_ft
    buf = buffer_space(55, jurisdiction="CDOT")
    assert not [
        p for p in placements if p.device_type == DeviceType.CONE and wz < p.station_ft < wz + buf
    ]


def _scenario() -> dict:
    return {
        "kind": "shoulder",
        "meta": {"project": "T", "address": "", "lat": 0.0, "lng": 0.0},
        "roadType": "rural_divided",
        "speed": 55,
        "lanes": 2,
        "laneWidth": 12.0,
        "divided": True,
        "workType": "utility_locate",
        "duration": "short",
        "workLen": 800.0,
        "night": False,
    }


def test_every_surface_counts_and_names_the_buffer_cones() -> None:
    c = TestClient(render_api.app)
    s = _scenario()
    # 495 ft buffer at 55 mph: ceil(495 / 110) = 5 intervals of 99 ft -> 4 cones.
    a = c.post("/render/audit", headers=AUTH, json=s).json()
    sp = a["sections"]["spacing"]
    assert sp["n_buffer_cones"] == 4
    assert sp["buffer_count_text"].startswith(
        "Buffer: 495 ft / 110 ft max spacing -> 4 cones at 99.0 ft intervals"
    )
    total = sp["n_cones_placed_total"]
    assert (
        f"9 tangent + 4 buffer + 2 downstream = {total} cones total" in sp["downstream_count_text"]
    )

    xl = load_workbook(io.BytesIO(c.post("/render/xlsx", headers=AUTH, json=s).content))
    xl_cones = sum(
        r[5]
        for r in xl["Device List"].iter_rows(min_row=2, values_only=True)
        if r[2] and "Cone" in str(r[2])
    )
    assert xl_cones == total

    md = c.post("/render/markdown", headers=AUTH, json=s).text
    assert "Then continue the line through the buffer space: 4 more traffic cones at 99 ft" in md

    pdf = pdfium.PdfDocument(c.post("/render/pdf", headers=AUTH, json=s).content)
    page1 = pdf[0].get_textpage().get_text_range()
    assert "BUFFER = 495 ft (NTS) · 4 CONES" in page1
    summary = re.search(r"Traffic Cone \(\d+-inch\) (\d+)", page1)
    assert summary is not None and int(summary.group(1)) == total
