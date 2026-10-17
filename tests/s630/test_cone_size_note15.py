"""#336 — the cone size follows CDOT S-630-1 Sheet 2 Note 15 on every surface.

Note 15: cones "shall be at least 28 inches"; the minimum is 36 inches
"on freeways and expressways, or during nighttime working hours", and
"should also be 36 inches" at 45 mph or more (ADT not collected; #336 S2
keeps the 45 mph clause at every ADT).

Rule 11: the size is asserted where it is printed — page 1 (``/render/pdf``),
the XLSX device list, the quote, the crew sheet (``/render/markdown`` and ``/render/crew-pdf``) and
the UI device breakdown — through the real API, plus the helper's matrix.
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.rules.devices import cone_display_name

from .test_cross_surface import (  # noqa: F401  (``api`` is a fixture)
    _pdf_text,
    _post,
    _stub_scan,
    _worst_case,
    api,
)

_36 = "Traffic Cone (36-inch)"
_28 = "Traffic Cone (28-inch)"


@pytest.mark.parametrize(
    ("speed", "night", "road_type", "expected"),
    [
        (25, False, "urban_low", _28),
        (40, False, "rural", _28),
        (44, False, "urban_high", _28),
        (25, True, "urban_low", _36),
        (40, True, "rural", _36),
        (40, False, "freeway", _36),
        (40, False, "expressway", _36),
        (45, False, "urban_high", _36),
        (65, False, "rural", _36),
    ],
)
def test_cone_display_name_note15_matrix(
    speed: int, night: bool, road_type: str, expected: str
) -> None:
    assert cone_display_name(speed, is_night=night, road_type=road_type) == expected


def _xlsx_rows(content: bytes, sheet: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(content), read_only=True)
    rows = list(wb[sheet].iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows


def _every_surface_cone_label(client: TestClient, scenario: dict[str, Any]) -> dict[str, list[str]]:
    """Each surface's cone label(s), read off the delivered bytes."""
    page1 = _pdf_text(_post(client, "/render/pdf", scenario).content)[0]
    crew = _post(client, "/render/markdown", scenario).text
    crew_pdf = "\n".join(_pdf_text(_post(client, "/render/crew-pdf", scenario).content))
    breakdown = _post(client, "/render/device-breakdown", scenario).json()["devices"]
    xlsx = _xlsx_rows(_post(client, "/render/xlsx", scenario).content, "Device List")
    quote_xlsx = _post(client, "/render/quote", {"scenario": scenario}).content
    quote = _xlsx_rows(quote_xlsx, "Equipment Detail")
    return {
        "page1": [s for s in (_36, _28) if s in page1],
        "crew": [s for s in (_36, _28) if s in crew],
        "crew_pdf": [s for s in (_36, _28) if s in crew_pdf],
        "breakdown": [str(r["device"]) for r in breakdown if "Cone" in str(r["device"])],
        "xlsx": [str(r[2]) for r in xlsx if r[1] == "CONE"],
        "quote": [str(r[3]) for r in quote if isinstance(r[0], int) and r[1] == "CONE"],
    }


@pytest.mark.parametrize(
    ("name", "overrides", "expected"),
    [
        # The issue's repro: N Broadway SB, left shoulder, 30 mph, night.
        ("broadway-left", {}, _36),
        ("adv-flagger", {}, _36),  # flagger, 40 mph, night
        ("adv-near-intersection", {}, _36),  # near intersection, 40 mph, night
        # Day plans below 45 mph keep 28-inch cones.
        ("broadway-left", {"night": False}, _28),
        ("control-typical", {}, _28),  # shoulder, 35 mph, day
    ],
    ids=[
        "shoulder-30-night",
        "flagger-40-night",
        "ni-40-night",
        "shoulder-30-day",
        "control-35-day",
    ],
)
def test_cone_size_is_one_label_on_every_surface(
    name: str,
    overrides: dict[str, Any],
    expected: str,
    api: TestClient,  # noqa: F811
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    scenario = {**_worst_case(name), **overrides}
    _stub_scan(name, monkeypatch)
    labels = _every_surface_cone_label(api, scenario)
    for surface, found in labels.items():
        assert found, f"{name}: no cone label on {surface}"
        assert set(found) == {expected}, f"{name}: {surface} prints {found}, expected {expected}"
