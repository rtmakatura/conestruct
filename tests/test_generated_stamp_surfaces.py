"""#268 — the four deliverables print ONE generated stamp (payload half).

Rule 11: test where the bug lives.  The pure helper is proven in
``test_generated_stamp.py``; this file drives the REAL render API with
``validators._utcnow`` pinned to one instant and reads the stamp back
out of each produced file:

  XLSX Summary  "Generated" row → a real datetime cell, yyyy-mm-dd
  quote         A3 "Generated" label + B3 datetime cell, yyyy-mm-dd
  crew markdown "**Generated:** YYYY-MM-DD" (base.md.j2)
  crew PDF      the same line, through the document renderer
  plan sheet    the title block's "DATE:" row on page 1

All five print the same UTC date.  The pinned instant is 2026-09-08
16:07:35 UTC — a date the pre-#268 code (three server-local clocks that
never read the seam) cannot produce on any day after it, so the proof
is red before the surfaces switch and cannot pass by coincidence.
"""

from __future__ import annotations

import io
import os
import re
from datetime import UTC, datetime

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

os.environ.setdefault("RENDER_API_SECRET", "test-secret-do-not-deploy")

from src.api.render_api import app  # noqa: E402
from src.rules import validators  # noqa: E402

HEADERS = {"Authorization": "Bearer test-secret-do-not-deploy"}
INSTANT = datetime(2026, 9, 8, 16, 7, 35, tzinfo=UTC)
STAMP = "2026-09-08"
DATE_FORMAT = "yyyy-mm-dd"


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture(autouse=True)
def _pin_clock(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(validators, "_utcnow", lambda: INSTANT)


def _shoulder_body() -> dict:
    return {
        "kind": "shoulder",
        "meta": {"project": "Stamp Test"},
        "roadType": "freeway",
        "speed": 55,
        "lanes": 2,
        "laneWidth": 12,
        "divided": True,
        "workType": "utility_locate",
        "duration": "short",
        "workLen": 1000,
        "night": False,
    }


def _pdf_text(body: bytes, page_index: int = 0) -> str:
    doc = pdfium.PdfDocument(body)
    try:
        return doc[page_index].get_textpage().get_text_bounded()
    finally:
        doc.close()


def test_xlsx_summary_generated_is_a_date_cell_on_the_pinned_instant(client: TestClient) -> None:
    res = client.post("/render/xlsx", headers=HEADERS, json=_shoulder_body())
    assert res.status_code == 200, res.text
    ws = load_workbook(io.BytesIO(res.content))["Summary"]
    rows = {
        ws.cell(row=r, column=1).value: ws.cell(row=r, column=2) for r in range(1, ws.max_row + 1)
    }
    cell = rows["Generated"]
    assert isinstance(cell.value, datetime), (
        f"Generated is {type(cell.value).__name__}, not a date cell"
    )
    assert cell.value == datetime(2026, 9, 8, 16, 7, 35)
    assert cell.number_format == DATE_FORMAT
    assert cell.value.strftime("%Y-%m-%d") == STAMP


def test_quote_row_3_is_a_generated_label_and_a_date_cell(client: TestClient) -> None:
    res = client.post("/render/quote", headers=HEADERS, json={"scenario": _shoulder_body()})
    assert res.status_code == 200, res.text
    ws = load_workbook(io.BytesIO(res.content))["Quote Summary"]
    assert ws["A3"].value == "Generated"
    assert isinstance(ws["B3"].value, datetime), f"B3 is {type(ws['B3'].value).__name__}"
    assert ws["B3"].value == datetime(2026, 9, 8, 16, 7, 35)
    assert ws["B3"].number_format == DATE_FORMAT


def test_crew_markdown_generated_line_prints_the_stamp(client: TestClient) -> None:
    res = client.post("/render/markdown", headers=HEADERS, json=_shoulder_body())
    assert res.status_code == 200, res.text
    m = re.search(r"\*\*Generated:\*\* (\S+)", res.text)
    assert m, "no Generated line in the crew narrative"
    assert m.group(1) == STAMP


def test_crew_pdf_generated_line_prints_the_stamp(client: TestClient) -> None:
    res = client.post("/render/crew-pdf", headers=HEADERS, json=_shoulder_body())
    assert res.status_code == 200, res.text
    text = _pdf_text(res.content)
    m = re.search(r"Generated:\s*(\d{4}-\d{2}-\d{2})", text)
    assert m, "no Generated line on the crew PDF's first page"
    assert m.group(1) == STAMP


def test_plan_sheet_date_prints_the_stamp(client: TestClient) -> None:
    res = client.post("/render/pdf", headers=HEADERS, json=_shoulder_body())
    assert res.status_code == 200, res.text
    text = _pdf_text(res.content)
    m = re.search(r"DATE:\s*(\d{4}-\d{2}-\d{2})", text)
    assert m, "no DATE row in the title block"
    assert m.group(1) == STAMP


def test_all_five_surfaces_print_the_same_date(client: TestClient) -> None:
    body = _shoulder_body()
    xlsx = load_workbook(
        io.BytesIO(client.post("/render/xlsx", headers=HEADERS, json=body).content)
    )["Summary"]
    xlsx_stamp = next(
        xlsx.cell(row=r, column=2).value
        for r in range(1, xlsx.max_row + 1)
        if xlsx.cell(row=r, column=1).value == "Generated"
    )
    quote = load_workbook(
        io.BytesIO(client.post("/render/quote", headers=HEADERS, json={"scenario": body}).content)
    )["Quote Summary"]
    md = client.post("/render/markdown", headers=HEADERS, json=body).text
    crew_pdf = _pdf_text(client.post("/render/crew-pdf", headers=HEADERS, json=body).content)
    plan = _pdf_text(client.post("/render/pdf", headers=HEADERS, json=body).content)

    stamps = {
        "xlsx": xlsx_stamp.strftime("%Y-%m-%d") if isinstance(xlsx_stamp, datetime) else xlsx_stamp,
        "quote": quote["B3"].value.strftime("%Y-%m-%d")
        if isinstance(quote["B3"].value, datetime)
        else quote["B3"].value,
        "markdown": re.search(r"\*\*Generated:\*\* (\S+)", md).group(1),  # type: ignore[union-attr]
        "crew_pdf": re.search(r"Generated:\s*(\d{4}-\d{2}-\d{2})", crew_pdf).group(1),  # type: ignore[union-attr]
        "plan": re.search(r"DATE:\s*(\d{4}-\d{2}-\d{2})", plan).group(1),  # type: ignore[union-attr]
    }
    assert set(stamps.values()) == {STAMP}, stamps
