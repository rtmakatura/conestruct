"""#309 R112/R114 — a lane closure near an intersection on a one-way street.

S-630-1 (July 2026) Sheet 2, General Note 8: "All warning and regulatory
signs shall be posted on both sides of the roadway on divided highways,
multi-lane ramps, one-way streets, and as directed by the Engineer, except
where only one shoulder is closed (ex: Case 11 on Sheet 7)."  A lane
closure is not a single-shoulder closure, so on a one-way street every
mainline sign is owed on the left curb too.

R114: Q1 the left-curb signs at -4.0 ft, CHOSEN; Q2 all eight mainline
labels mirrored, cross-street sets unchanged; Q3 (a) a ``divided`` verdict
keeps today's plan; Q4 ``undecided`` is refused as the shoulder kind's is;
Q5 the one-way wording.  "A two-way near-intersection plan stays
byte-identical" (R112).

Tested at the payload chokepoint (``render_api._placements_for``), on the
endpoints, and on the rendered crew sheet and plan page (Rule 11).  The
Broadway body is the checkpoint's repro (validation-artifacts/committed/
issue-309-oneway-near-intersection/probes/ni_broadway_probe.py).
"""

from __future__ import annotations

import copy
import datetime
import io
import os
from typing import Any

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import TypeAdapter

from src.api import render_api
from src.api.render_api import app
from src.api.schemas import Scenario
from src.narrative.crew_narrative import render_crew_narrative_markdown
from src.rendering import plan_sheet as ps
from src.rules import site_detection as sd
from src.rules.devices import DeviceType
from src.rules.validators import scenario_display_name

AUTH = {"Authorization": "Bearer test-secret-do-not-deploy"}
ONE_WAY = {
    "oneway": "yes",
    "highwayClass": "primary",
    "twinDistanceM": None,
    "twinSearched": True,
    "confirmed": None,
}
MAINLINE_LABELS = ("W4-2R", "W20-5R", "W20-1", "R2-10", "G20-5P", "G20-1", "G20-2", "R2-11")


@pytest.fixture(scope="module", autouse=True)
def _env() -> None:
    os.environ["RENDER_API_SECRET"] = "test-secret-do-not-deploy"


@pytest.fixture(autouse=True)
def _no_overpass(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        sd, "_overpass_request_with_fallback", lambda *_a, **_kw: ({"elements": []}, None)
    )


def _body(carriageway: dict[str, Any] | None = None, **over: Any) -> dict[str, Any]:
    """N Broadway SB at E 11th Ave (39.73370, -104.98753), corridor_end."""
    b: dict[str, Any] = {
        "kind": "near_intersection",
        "meta": {
            "project": "309",
            "address": "N Broadway SB at E 11th Ave",
            "lat": 39.73370,
            "lng": -104.98753,
            "bearingDeg": 0.0,
        },
        "roadType": "urban_arterial",
        "speed": 30,
        "lanes": 4,
        "laneWidth": 10.5,
        "divided": False,
        "workType": "utility_cut",
        "duration": "short",
        "workLen": 500.0,
        "night": False,
        "approaches": [
            {
                "id": "e11_wb",
                "speed": 25,
                "roadType": "urban_arterial",
                "lanesPerDirection": 1,
                "laneWidth": 11.0,
                "signalized": True,
                "alongStationFt": -100.0,
            },
            {
                "id": "e11_eb",
                "speed": 25,
                "roadType": "urban_arterial",
                "lanesPerDirection": 1,
                "laneWidth": 11.0,
                "signalized": True,
                "alongStationFt": -100.0,
            },
        ],
    }
    if carriageway is not None:
        b["carriageway"] = carriageway
    b.update(over)
    return b


def _plan(body: dict[str, Any]):
    placements, params, *_ = render_api._placements_for(
        TypeAdapter(Scenario).validate_python(copy.deepcopy(body))
    )
    return placements, params


def _mainline_signs(placements) -> list:
    return [
        p
        for p in placements
        if p.device_type == DeviceType.SIGN_GENERIC and p.approach_id == "mainline"
    ]


client = TestClient(app)


def _post(path: str, body: dict[str, Any]):
    return client.post(path, headers=AUTH, json=body)


def _note8(audit: dict[str, Any]) -> dict[str, Any]:
    return next(c for c in audit["sections"]["colorado"]["checks"] if "both sides" in c["label"])


# --- Surface C: the relay reaches the plan -------------------------------------


def test_the_one_way_fact_reaches_the_plan() -> None:
    _, params = _plan(_body(ONE_WAY))
    assert params.one_way_street is True
    assert params.is_divided is False


# --- Surface A: the mirror ------------------------------------------------------


def test_every_mainline_sign_is_mirrored_to_the_left_curb_at_minus_4_ft() -> None:
    placements, _ = _plan(_body(ONE_WAY))
    signs = _mainline_signs(placements)
    right = sorted((p.label, p.station_ft) for p in signs if p.offset_ft > 0)
    left = sorted((p.label, p.station_ft) for p in signs if p.offset_ft < 0)
    assert right == left
    assert {label for label, _ in left} == set(MAINLINE_LABELS)
    assert all(p.offset_ft == -4.0 for p in signs if p.offset_ft < 0)
    assert len(left) == 8 and len(placements) == 42


def test_the_cross_street_sets_are_unchanged() -> None:
    def approach_signs(body):
        placements, _ = _plan(body)
        return sorted(
            (p.approach_id, p.label, p.station_ft, p.offset_ft)
            for p in placements
            if p.approach_id != "mainline"
        )

    assert approach_signs(_body(ONE_WAY)) == approach_signs(_body())


def test_the_left_curb_signs_land_on_the_drawn_left_strip() -> None:
    """Not R94's float: offset 0 is the left lane edge on a one-way street,
    and the left curb strip is drawn from there to -shoulder."""
    placements, params = _plan(_body(ONE_WAY))
    top, _ = ps._road_y_extent(params, params.shoulder_width_ft)
    left = [p for p in _mainline_signs(placements) if p.offset_ft < 0]
    assert left
    for p in left:
        assert ps.PLAN_Y_CENTER <= ps._y_of(p.offset_ft, False) <= top


# --- Surface B: the check, the audit, the title, the crew sheet -----------------


def test_the_note_8_row_is_required_and_met() -> None:
    r = _post("/render/audit", _body(ONE_WAY))
    assert r.status_code == 200, r.text
    row = _note8(r.json())
    assert row["pass"] is True
    assert row["label"] == "Signs on both sides of one-way street"
    assert row["detail"].startswith("Required: True. Signs placed: 8 left, 8 right.")
    # The SHOULDER WORK figure note is about a shoulder closure (R90).
    assert "Fig. 6P-3" not in row["detail"]


def test_the_narrative_states_what_the_plan_does() -> None:
    case = _post("/render/audit", _body(ONE_WAY)).json()["sections"]["case"]
    assert (
        "applied to a one-way street with every mainline warning and regulatory "
        "sign posted on both sides of the roadway per CDOT S-630-1 Sheet 2 General "
        "Note 8" in case["narrative"]
    )
    assert "undivided highway" not in case["narrative"]
    assert (
        "(2) A one-way street has no opposing mainline direction; the plate's "
        "opposing-direction signing does not apply." in case["narrative_2"]
    )
    assert "opposing mainline direction is not signed" not in case["narrative_2"]


def test_the_title_names_the_one_way_street() -> None:
    _, params = _plan(_body(ONE_WAY))
    assert scenario_display_name(params) == "Lane Closure Near Intersection · One-Way Street"


def _crew(body: dict[str, Any]) -> str:
    from src.api.schemas import scenario_to_call

    scenario = TypeAdapter(Scenario).validate_python(copy.deepcopy(body))
    placements, params, *_ = render_api._placements_for(scenario)
    _, _, kwargs = scenario_to_call(scenario)
    return render_crew_narrative_markdown(placements, params, approaches=kwargs["approaches"])


def test_the_crew_sheet_says_both_sides() -> None:
    md = _crew(_body(ONE_WAY))
    assert (
        "One-way street: all warning and regulatory signs are posted on BOTH sides "
        "of the roadway per CDOT S-630-1 Sheet 2 General Note 8." in md
    )
    assert "closed shoulder's side only" not in md


def test_the_crew_steps_place_the_left_curb_signs() -> None:
    """Every crew step that places a mainline sign says where its left-curb
    twin goes: the G20-2 (step 2), the advance series (step 6), and the
    regulatory and bookend signs (step 7, no longer "single side")."""
    md = _crew(_body(ONE_WAY))
    assert "Place matching sign on the left curb, 4 ft left of the left edge of the roadway." in md
    assert (
        "Place matching signs on the left curb at each location, "
        "4 ft left of the left edge of the roadway" in md
    )
    assert "single side" not in md
    assert (
        "on both sides: the right side 46 ft from the left edge of the roadway, "
        "the left curb 4 ft left of it" in md
    )


def test_a_two_way_crew_sheet_is_unchanged() -> None:
    assert _crew(_body(TWO_WAY)) == _crew(_body())
    assert "single side, 46 ft from centerline" in _crew(_body())


def test_the_plan_sheet_s_fine_print_states_departure_2_the_same_way() -> None:
    """Q5's departure (2) on page 1's notes, as in the audit."""
    page1 = " ".join(_pdf_pages(_body(ONE_WAY))[0].split())
    assert "A one-way street has no opposing mainline direction" in page1
    assert "Opposing mainline direction not signed" not in page1


def test_the_device_breakdown_counts_both_curbs() -> None:
    def qty(body, code):
        rows = _post("/render/device-breakdown", body).json()["devices"]
        return sum(d["qty"] for d in rows if d.get("code") == code)

    # Mainline-only codes: one on the two-way plan, one per curb here.
    for code in ("W4-2R", "W20-5R", "G20-1", "G20-2"):
        assert qty(_body(), code) == 1, code
        assert qty(_body(ONE_WAY), code) == 2, code


def test_the_pdf_and_xlsx_render() -> None:
    for path in ("/render/pdf", "/render/xlsx"):
        assert _post(path, _body(ONE_WAY)).status_code == 200, path


# --- What does not change -------------------------------------------------------


TWO_WAY = {**ONE_WAY, "oneway": "no", "highwayClass": "residential"}


@pytest.mark.parametrize("path", ["/render/audit", "/render/device-breakdown"])
def test_a_two_way_plan_is_byte_identical_with_or_without_the_facts(path: str) -> None:
    """R112: "A two-way near-intersection plan stays byte-identical."  The
    frontend now relays the facts on every road; a two-way road's are
    ``not_applicable`` and change nothing."""
    assert _post(path, _body(TWO_WAY)).content == _post(path, _body()).content


def _xlsx_cells(body: dict[str, Any]) -> list[tuple[str, str, Any]]:
    """Every cell but the generated stamp (a datetime that differs on every
    render of the same body, so raw bytes can't compare)."""
    wb = load_workbook(io.BytesIO(_post("/render/xlsx", body).content))
    return [
        (ws.title, c.coordinate, c.value)
        for ws in wb
        for row in ws.iter_rows()
        for c in row
        if c.value is not None and not isinstance(c.value, datetime.datetime)
    ]


def _pdf_pages(body: dict[str, Any]) -> list[str]:
    """Each page's text (the PDF's metadata carries a creation time)."""
    doc = pdfium.PdfDocument(io.BytesIO(_post("/render/pdf", body).content))
    return [doc[i].get_textpage().get_text_range() for i in range(len(doc))]


def test_a_two_way_plan_s_xlsx_and_pdf_carry_the_same_content() -> None:
    assert _xlsx_cells(_body(TWO_WAY)) == _xlsx_cells(_body())
    assert _pdf_pages(_body(TWO_WAY)) == _pdf_pages(_body())


def test_a_divided_verdict_keeps_today_s_plan() -> None:
    """R114 Q3 (a): divided near-intersection plans are a follow-up issue."""
    divided = {**ONE_WAY, "twinDistanceM": 14.8}
    placements, params = _plan(_body(divided))
    assert params.one_way_street is False and params.is_divided is False
    assert _post("/render/audit", _body(divided)).content == _post("/render/audit", _body()).content
    assert len(placements) == 34


def test_an_undecided_road_is_refused_honestly() -> None:
    """R114 Q4: the shoulder kind's treatment, a 400 the WHAT band's
    carriageway row answers."""
    r = _post("/render/audit", _body({**ONE_WAY, "twinSearched": False}))
    assert r.status_code == 400
    assert r.json()["detail"]["error"] == "carriageway_undecided"


def test_the_operator_s_answer_decides_an_undecided_road() -> None:
    answered = {**ONE_WAY, "twinSearched": False, "confirmed": "one_way_street"}
    _, params = _plan(_body(answered))
    assert params.one_way_street is True
