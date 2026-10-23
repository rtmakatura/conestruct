"""#372: page 1's device summary counts the cross-street devices.

A near-intersection plan draws the mainline only (#117 Option C), but its
TRAFFIC CONTROL DEVICE SUMMARY gives quantities and a total, so it must
equal the device XLSX, cross-street signs included (P2, one voice per
fact).  Rule-11 layer: the text of the PDF and the cells of the XLSX the
real endpoints return.
"""

from __future__ import annotations

import copy
import io
from typing import Any

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.render_api import app
from tests.test_plan_sheet_device_summary import FLAGGER_BODY, SHOULDER_BODY

AUTH = {"Authorization": "Bearer test-secret"}


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setenv("RENDER_API_SECRET", "test-secret")


@pytest.fixture()
def client():
    return TestClient(app)


def _approach(approach_id: str, **overrides: Any) -> dict[str, Any]:
    return {
        "id": approach_id,
        "speed": 30,
        "roadType": "urban_arterial",
        "lanesPerDirection": 1,
        "laneWidth": 12.0,
        "signalized": True,
        **overrides,
    }


def _ni_body(approaches: list[dict[str, Any]], **overrides: Any) -> dict[str, Any]:
    # meta.lat = lng = 0 keeps the site scan off the network.
    return {
        "kind": "near_intersection",
        "meta": {"project": "issue 372", "address": "", "lat": 0.0, "lng": 0.0},
        "roadType": "urban_arterial",
        "speed": 30,
        "lanes": 3,
        "laneWidth": 12.0,
        "divided": False,
        "workType": "utility_cut",
        "duration": "short",
        "workLen": 500.0,
        "night": False,
        "approaches": approaches,
        **overrides,
    }


NI_BODIES = {
    # N1's shape: both cross-street legs, Denver arterial.
    "two-legs-denver": _ni_body(
        [_approach("cross_a", alongStationFt=-100.0), _approach("cross_b", alongStationFt=-100.0)],
        jurisdiction_key="denver",
        street_class="arterial",
    ),
    # N6's shape: one leg.
    "one-leg": _ni_body([_approach("cross_a", alongStationFt=-100.0)]),
    "two-legs-no-jurisdiction": _ni_body(
        [_approach("cross_a", alongStationFt=-100.0), _approach("cross_b", alongStationFt=-100.0)]
    ),
}


def _summary(client: TestClient, body: dict, tmp_path) -> tuple[list[tuple[str, int]], int]:
    """(name, qty) rows and the total from page 1's device summary."""
    resp = client.post("/render/pdf", json=body, headers=AUTH)
    assert resp.status_code == 200, resp.text
    path = tmp_path / "sheet.pdf"
    path.write_bytes(resp.content)
    pdf = pdfium.PdfDocument(str(path))
    try:
        lines = pdf[0].get_textpage().get_text_range().splitlines()
    finally:
        pdf.close()
    start = lines.index("TRAFFIC CONTROL DEVICE SUMMARY") + 2  # past the column header
    end = next(i for i in range(start, len(lines)) if lines[i].startswith("TOTAL DEVICES"))
    rows = []
    for line in lines[start:end]:
        name, qty = line.rsplit(" ", 1)
        rows.append((name.removeprefix("— "), int(qty)))
    return rows, int(lines[end].removeprefix("TOTAL DEVICES "))


def _xlsx(client: TestClient, body: dict) -> tuple[list[tuple[str, int]], int]:
    """(description, qty) rows and the total device count from the device XLSX."""
    resp = client.post("/render/xlsx", json=body, headers=AUTH)
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    try:
        rows = [
            (str(r[2]), int(r[5])) for r in wb["Device List"].iter_rows(min_row=2, values_only=True)
        ]
        summary = dict(wb["Summary"].iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    return rows, int(summary["Total device count"])


def _assert_sheet_equals_xlsx(client: TestClient, body: dict, tmp_path) -> int:
    sheet_rows, sheet_total = _summary(client, body, tmp_path)
    xlsx_rows, xlsx_total = _xlsx(client, body)
    assert sheet_total == xlsx_total
    assert [q for _, q in sheet_rows] == [q for _, q in xlsx_rows]
    for (sheet_name, _), (xlsx_name, _) in zip(sheet_rows, xlsx_rows, strict=True):
        # The sheet truncates a long name to its column; the start must match.
        assert xlsx_name.startswith(sheet_name.rstrip("…").rstrip(".")), (sheet_name, xlsx_name)
    return sheet_total


@pytest.mark.parametrize("case", sorted(NI_BODIES))
def test_ni_page1_summary_equals_the_device_xlsx(client, tmp_path, case: str) -> None:
    _assert_sheet_equals_xlsx(client, copy.deepcopy(NI_BODIES[case]), tmp_path)


def test_ni_page1_counts_each_cross_street_leg(client, tmp_path) -> None:
    """Each leg carries its own W20-1, R2-10, R2-11 (S-630-1 Sheet 10), so
    two legs count more of each than one; the mainline-only count would be
    the same for both."""

    def counts(body: dict) -> dict[str, int]:
        rows, _ = _summary(client, body, tmp_path)
        return {name.split(" ", 1)[0]: q for name, q in rows if name[:1] in "WR"}

    one = counts(copy.deepcopy(NI_BODIES["one-leg"]))
    two = counts(copy.deepcopy(NI_BODIES["two-legs-no-jurisdiction"]))
    for code in ("W20-1", "R2-10", "R2-11"):
        assert two[code] > one[code] > 1, (code, one, two)


@pytest.mark.parametrize("body", [FLAGGER_BODY, SHOULDER_BODY], ids=["flagger", "shoulder"])
def test_other_kinds_still_equal_the_device_xlsx(client, tmp_path, body: dict) -> None:
    """Guard: the kinds that draw every placement were already equal."""
    _assert_sheet_equals_xlsx(client, copy.deepcopy(body), tmp_path)
