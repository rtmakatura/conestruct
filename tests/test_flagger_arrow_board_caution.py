"""#379: a flagger plan's arrow board runs in caution mode only, on every surface.

MUTCD 11th Ed. §6L.06 ¶18 (Standard), p. 833: "for temporarily closing one
lane on a two-lane, two-way roadway, an arrow board shall be used only in the
caution mode."  Denver's ">=1 lane" rule (and Englewood's, Greeley's and
Lakewood's on their classes) adds a board to a flagger plan that the layout
doesn't place.  Rule 11: every assertion reads a real render endpoint's
output -- page 1, the device XLSX, the quote XLSX, the crew sheet and the
breakdown the audit carries -- not the helper.
"""

from __future__ import annotations

import copy
import io
import re

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.render_api import app
from tests.test_plan_sheet_device_summary import FLAGGER_BODY, SHOULDER_BODY
from tests.test_r117_jurisdiction_raised import _broadway_denver

AUTH = {"Authorization": "Bearer test-secret"}
NAME = "Arrow Board (caution mode only)"
CITE = "§6L.06 ¶18"


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setenv("RENDER_API_SECRET", "test-secret")


@pytest.fixture()
def client():
    return TestClient(app)


def _flagger(key: str = "denver", street_class: str | None = None) -> dict:
    body = copy.deepcopy(FLAGGER_BODY)
    body["jurisdiction_key"] = key
    if street_class is not None:
        body["street_class"] = street_class
    return body


def _post(client: TestClient, path: str, body: dict) -> bytes:
    resp = client.post(path, json=body, headers=AUTH)
    assert resp.status_code == 200, resp.text
    return resp.content


def _page1(client: TestClient, body: dict, tmp_path) -> str:
    tmp = tmp_path / "sheet.pdf"
    tmp.write_bytes(_post(client, "/render/pdf", body))
    pdf = pdfium.PdfDocument(str(tmp))
    try:
        return pdf[0].get_textpage().get_text_range()
    finally:
        pdf.close()


def _sheet_rows(content: bytes, sheet: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(content))
    try:
        return [r for r in wb[sheet].iter_rows(min_row=2, values_only=True) if r and r[1]]
    finally:
        wb.close()


def test_page_1_names_the_board_caution_mode_only(client, tmp_path) -> None:
    text = _page1(client, _flagger(), tmp_path)
    assert re.search(r"Arrow Board \(caution\s+mode only\)", text), text


def test_device_xlsx_row_names_the_mode_and_cites_6l06(client) -> None:
    rows = _sheet_rows(_post(client, "/render/xlsx", _flagger()), "Device List")
    (arrow,) = [r for r in rows if "Arrow Board" in str(r[1])]
    assert arrow[1] == NAME
    assert arrow[5] == 1
    assert "Jurisdiction-required: DOTI PT-116.1." in arrow[6]
    assert CITE in arrow[6] and "never an arrow or chevron" in arrow[6]


def test_quote_xlsx_row_names_the_mode_and_cites_6l06(client) -> None:
    content = _post(client, "/render/quote", {"scenario": _flagger()})
    rows = [r for r in _sheet_rows(content, "Equipment Detail") if isinstance(r[0], int)]
    (arrow,) = [r for r in rows if "Arrow Board" in str(r[1])]
    assert arrow[1] == NAME
    assert CITE in arrow[9]


def test_crew_sheet_lists_and_places_the_board_in_caution_mode(client) -> None:
    md = _post(client, "/render/markdown", _flagger()).decode("utf-8")
    equipment = md.split("## Required Equipment", 1)[1].split("\n## ", 1)[0]
    assert f"1× {NAME}" in equipment
    setup = md.split("## Setup Procedure", 1)[1].split("## Takedown Procedure", 1)[0]
    assert "Arrow board (DOTI PT-116.1 requires one): Caution mode only" in setup
    assert CITE in setup and "The Traffic Control Supervisor picks its position." in setup
    takedown = md.split("## Takedown Procedure", 1)[1]
    assert "6. Deactivate and remove the arrow board." in takedown
    # No surface states or implies arrow or chevron mode.
    assert "ARROW mode" not in md


def test_crew_pdf_carries_the_same_lines(client, tmp_path) -> None:
    tmp = tmp_path / "crew.pdf"
    tmp.write_bytes(_post(client, "/render/crew-pdf", _flagger()))
    pdf = pdfium.PdfDocument(str(tmp))
    try:
        text = "\n".join(p.get_textpage().get_text_range() for p in pdf)
    finally:
        pdf.close()
    assert "caution mode only" in text and "Deactivate and remove the arrow board" in text


def test_breakdown_and_its_rule_say_caution_mode(client) -> None:
    body = client.post("/render/device-breakdown", headers=AUTH, json=_flagger()).json()
    (row,) = [d for d in body["devices"] if "Arrow Board" in d["device"]]
    assert row["device"] == NAME
    (rule,) = [
        d
        for d in body["jurisdiction"]["applied_deltas"]
        if d["effect"].get("device") == "arrow_board"
    ]
    assert rule["device_label"] == "Arrow board (caution mode only)"
    assert CITE in rule["mode_note"]
    assert rule["raised"] is True


@pytest.mark.parametrize(
    ("key", "street_class", "name"),
    [
        ("englewood", "collector", NAME),
        # The record's own name keeps its type: "Arrow Board (Type C)".
        ("greeley", "arterial", "Arrow Board (Type C, caution mode only)"),
        ("lakewood", "arterial", NAME),
    ],
)
def test_every_jurisdiction_board_on_a_flagger_plan_is_caution_only(
    client, key, street_class, name
) -> None:
    rows = _sheet_rows(_post(client, "/render/xlsx", _flagger(key, street_class)), "Device List")
    (arrow,) = [r for r in rows if "Arrow" in str(r[1])]
    assert arrow[1] == name and CITE in arrow[6]


def test_near_intersection_board_is_unchanged(client) -> None:
    body = client.post("/render/device-breakdown", headers=AUTH, json=_broadway_denver()).json()
    (rule,) = [
        d
        for d in body["jurisdiction"]["applied_deltas"]
        if d["effect"].get("device") == "arrow_board"
    ]
    assert rule["device_label"] == "Arrow board" and "mode_note" not in rule
    assert all("caution mode only" not in d["device"] for d in body["devices"])


def test_shoulder_plan_is_unchanged(client) -> None:
    body = copy.deepcopy(SHOULDER_BODY)
    body["jurisdiction_key"] = "denver"
    md = _post(client, "/render/markdown", body).decode("utf-8")
    assert NAME not in md and "requires one): Caution" not in md
    rows = _sheet_rows(_post(client, "/render/xlsx", body), "Device List")
    assert all("caution mode only" not in str(r[1]) for r in rows)


def test_flagger_without_a_jurisdiction_has_no_board(client) -> None:
    md = _post(client, "/render/markdown", copy.deepcopy(FLAGGER_BODY)).decode("utf-8")
    assert "arrow board" not in md.lower()


def test_audit_breakdown_says_caution_mode(client) -> None:
    body = _flagger()
    body["include_breakdown"] = True
    audit = client.post("/render/audit", headers=AUTH, json=body).json()
    (row,) = [d for d in audit["breakdown"]["devices"] if "Arrow Board" in d["device"]]
    assert row["device"] == NAME
    (rule,) = [
        d
        for d in audit["breakdown"]["jurisdiction"]["applied_deltas"]
        if d["effect"].get("device") == "arrow_board"
    ]
    assert CITE in rule["mode_note"]
