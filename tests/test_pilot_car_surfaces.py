"""#373: a flagger plan's pilot car is listed on every surface that lists equipment.

The pilot car isn't a placement: its G20-4 rides the rear of the vehicle
(S-630-1 Sheet 26), so layout.py places no roadside sign for it.  Before #373
only the crew sheet and the audit's flagger line read the ``pilotCar`` fact;
page 1's summary, the device XLSX, the quote XLSX and the breakdown (the
screen panel and the audit's device list) left it out.  Rule 11: every
assertion reads a real render endpoint's output.
"""

from __future__ import annotations

import copy
import io
import json
import re

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.render_api import app
from tests.test_plan_sheet_device_summary import FLAGGER_BODY

AUTH = {"Authorization": "Bearer test-secret"}
NAME = 'Pilot car with G20-4 "PILOT CAR/FOLLOW ME" sign'
SHEET_26 = "S-630-1 Sheet 26"


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setenv("RENDER_API_SECRET", "test-secret")


@pytest.fixture()
def client():
    return TestClient(app)


def _flagger(pilot_car: bool = True) -> dict:
    body = copy.deepcopy(FLAGGER_BODY)
    body["pilotCar"] = pilot_car
    return body


def _post(client: TestClient, path: str, body: dict) -> bytes:
    resp = client.post(path, json=body, headers=AUTH)
    assert resp.status_code == 200, resp.text
    return resp.content


def _page1(client: TestClient, body: dict, tmp_path) -> str:
    tmp = tmp_path / "sheet.pdf"
    tmp.write_bytes(_post(client, "/render/pdf", body))
    pdf = pdfium.PdfDocument(str(tmp))
    try:
        return pdf[0].get_textpage().get_text_range()
    finally:
        pdf.close()


def _sheet_rows(content: bytes, sheet: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(content))
    try:
        return [r for r in wb[sheet].iter_rows(min_row=2, values_only=True) if r and r[1]]
    finally:
        wb.close()


def _summary_total(content: bytes) -> int:
    rows = _sheet_rows(content, "Summary")
    return next(int(r[1]) for r in rows if r[0] == "Total device count")


def _page1_total(text: str) -> int:
    m = re.search(r"TOTAL DEVICES\s*(\d+)", text)
    assert m, text
    return int(m.group(1))


def test_page_1_lists_one_pilot_car_and_counts_it(client, tmp_path) -> None:
    with_car = _page1(client, _flagger(), tmp_path)
    without = _page1(client, _flagger(False), tmp_path)
    # Page 1 truncates long names to the column; the G20-4 lead must fit.
    assert re.search(r"Pilot car with G20-4", with_car), with_car
    assert "Pilot car" not in without
    assert _page1_total(with_car) == _page1_total(without) + 1


def test_device_xlsx_lists_one_pilot_car_with_no_pay_item(client) -> None:
    content = _post(client, "/render/xlsx", _flagger())
    rows = [r for r in _sheet_rows(content, "Device List") if r[2] == NAME]
    assert len(rows) == 1, rows
    _item, type_label, _desc, pay_item, unit, qty, note = rows[0]
    assert (type_label, pay_item, unit, qty) == (NAME, "—", "EACH", 1)
    assert SHEET_26 in note
    without = _post(client, "/render/xlsx", _flagger(False))
    assert _summary_total(content) == _summary_total(without) + 1


def test_page_1_total_equals_the_device_xlsx_total(client, tmp_path) -> None:
    """#372's invariant holds with the pilot car on both surfaces."""
    body = _flagger()
    assert _page1_total(_page1(client, body, tmp_path)) == _summary_total(
        _post(client, "/render/xlsx", body)
    )


def test_quote_xlsx_lists_the_pilot_car_unpriced(client) -> None:
    with_car = _post(client, "/render/quote", {"scenario": _flagger()})
    rows = [r for r in _sheet_rows(with_car, "Equipment Detail") if r[1] == NAME]
    assert len(rows) == 1, rows
    _item, _dt, _label, _desc, qty, unit, rate, _days, extended, note = rows[0]
    assert (qty, unit, rate, extended) == (1, "EACH", "—", "—")
    assert SHEET_26 in note and "NOT included in the quote total" in note


def test_quote_total_is_unchanged_by_the_unpriced_line(client) -> None:
    def breakdown(pilot_car: bool) -> dict:
        resp = client.post(
            "/render/quote-breakdown", json={"scenario": _flagger(pilot_car)}, headers=AUTH
        )
        assert resp.status_code == 200, resp.text
        return resp.json()

    with_car, without = breakdown(True), breakdown(False)
    assert with_car["total"] == without["total"]
    lines = [ln for ln in with_car["equipment_lines"] if ln["device_type"] == NAME]
    assert len(lines) == 1 and lines[0]["unpriced"] is True
    # The key appears only on the unpriced line.
    assert not any("unpriced" in ln for ln in without["equipment_lines"])


def test_breakdown_and_audit_device_list_list_the_pilot_car(client) -> None:
    breakdown = json.loads(_post(client, "/render/device-breakdown", _flagger()))
    rows = [r for r in breakdown["devices"] if r["device"] == NAME]
    assert rows == [{"device": NAME, "code": "—", "function": "Traffic control", "qty": 1}]
    without = json.loads(_post(client, "/render/device-breakdown", _flagger(False)))
    assert breakdown["total_devices"] == without["total_devices"] + 1

    audit = json.loads(_post(client, "/render/audit", {**_flagger(), "include_breakdown": True}))
    assert audit["breakdown"] == breakdown
    assert audit["sections"]["flagger"]["pilot_car_text"] == "1 vehicle"


def test_crew_sheet_names_the_pilot_car_with_the_shared_name(client) -> None:
    md = _post(client, "/render/markdown", _flagger()).decode("utf-8")
    assert md.count(f"1× {NAME} mounted on the rear of the vehicle") == 1


@pytest.mark.parametrize(
    "path,wrap",
    [
        ("/render/xlsx", False),
        ("/render/device-breakdown", False),
        ("/render/quote-breakdown", True),
        ("/render/markdown", False),
    ],
)
def test_no_pilot_car_means_no_pilot_car_row(client, path, wrap) -> None:
    body = _flagger(False)
    content = _post(client, path, {"scenario": body} if wrap else body)
    if path.endswith("xlsx"):
        assert not any(
            "Pilot car" in str(c) for r in _sheet_rows(content, "Device List") for c in r
        )
    else:
        assert b"Pilot car" not in content
