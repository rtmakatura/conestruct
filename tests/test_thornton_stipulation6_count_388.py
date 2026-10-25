"""#388: Thornton's stipulation 6 counts two signs per direction of travel
affected (R1, Ryan 2026-10-24), and every surface prints the same count.

Stipulation 6, as the repo records it (``data/jurisdictions/thornton.json``):
"CONSTRUCTION AHEAD and MEN WORKING signs required for each direction of
travel affected (stipulation #6)."  A direction is affected when its traffic
passes the work on the same carriageway, read from the plan's road facts
(R7): 1 on a one-way street or one carriageway of a divided road, 2 on a
two-way undivided road (every flagger plan); a near_intersection plan adds
each cross-street leg it signs.  Rule 11: the assertions read the real endpoints' output
(device XLSX, quote XLSX, breakdown JSON, crew markdown, page 1 of the PDF).
"""

from __future__ import annotations

import copy
import io
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.render_api import _jurisdiction_eval, app
from src.rules.jurisdiction import (
    aggregate_device_rows_with_deltas,
    annotate_count_effects,
    load_jurisdiction,
)
from tests.test_plan_sheet_device_summary import (
    FLAGGER_BODY,
    SHOULDER_BODY,
    pdf_text,
)
from tests.test_r117_jurisdiction_raised import _broadway_denver

AUTH = {"Authorization": "Bearer test-secret"}
AWS = "Advance Warning Signs"


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setenv("RENDER_API_SECRET", "test-secret")


@pytest.fixture()
def client():
    return TestClient(app)


def _with(base: dict, key: str = "thornton", **extra) -> dict:
    body = copy.deepcopy(base)
    body["jurisdiction_key"] = key
    body.update(extra)
    return body


def _ni_two_legs() -> dict:
    """The Broadway body with a second cross-street leg."""
    body = _broadway_denver()
    leg = copy.deepcopy(body["approaches"][0])
    leg["id"] = "w11"
    body["approaches"].append(leg)
    return body


# The carriageway facts of a one-way street (#308): a street class tagged
# one-way, no same-name twin in the searched pool -> ``one_way_street``.
ONE_WAY = {"oneway": "yes", "highwayClass": "residential", "twinSearched": True}

UNDIVIDED_SHOULDER = {
    **SHOULDER_BODY,
    "roadType": "rural_undivided",
    "divided": False,
    "lanes": 1,
}
ONE_WAY_SHOULDER = {
    **SHOULDER_BODY,
    "roadType": "urban_arterial",
    "speed": 30,
    "divided": False,
    "carriageway": ONE_WAY,
}

# case -> (body, directions of travel affected, R7)
CASES = {
    "shoulder-divided": (lambda: _with(SHOULDER_BODY), 1),
    "shoulder-divided-night": (lambda: _with(SHOULDER_BODY, night=True), 1),
    "shoulder-two-way": (lambda: _with(UNDIVIDED_SHOULDER), 2),
    "shoulder-one-way": (lambda: _with(ONE_WAY_SHOULDER), 1),
    "flagger": (lambda: _with(FLAGGER_BODY), 2),
    "near_intersection-two-way-1-leg": (lambda: _with(_broadway_denver()), 3),
    "near_intersection-two-way-2-legs": (lambda: _with(_ni_two_legs()), 4),
    "near_intersection-one-way-1-leg": (
        lambda: _with(_broadway_denver(), carriageway=ONE_WAY),
        2,
    ),
}


def _post(client: TestClient, path: str, body: dict) -> bytes:
    resp = client.post(path, json=body, headers=AUTH)
    assert resp.status_code == 200, resp.text
    return resp.content


def _xlsx_qty(content: bytes, sheet: str, name_col: int, qty_col: int) -> list:
    wb = load_workbook(io.BytesIO(content))
    try:
        return [
            r[qty_col]
            for r in wb[sheet].iter_rows(min_row=2, values_only=True)
            if r and str(r[name_col] or "").startswith(AWS)
        ]
    finally:
        wb.close()


def _counts(client: TestClient, body: dict, tmp_path) -> dict[str, list]:
    xlsx = _post(client, "/render/xlsx", body)
    quote = client.post("/render/quote", json={"scenario": body}, headers=AUTH)
    assert quote.status_code == 200, quote.text
    breakdown = client.post("/render/device-breakdown", json=body, headers=AUTH).json()
    md = _post(client, "/render/markdown", body).decode()
    page = pdf_text(client, body, tmp_path)
    return {
        "device_xlsx": _xlsx_qty(xlsx, "Device List", 1, 5),
        "quote_xlsx": _xlsx_qty(quote.content, "Equipment Detail", 1, 4),
        "breakdown": [d["qty"] for d in breakdown["devices"] if str(d["device"]).startswith(AWS)],
        # What JurisdictionSection's deltaImpact shows (d.qty ?? effect.qty ?? 1).
        "delta_panel": [
            d.get("qty", (d.get("effect") or {}).get("qty", 1))
            for d in breakdown["jurisdiction"]["applied_deltas"]
            if (d.get("effect") or {}).get("device") == "advance_warning_signs"
        ],
        "crew_sheet": [int(n) for n in re.findall(rf"- (\d+)× {AWS}", md)],
        "page_1": [
            int(n) for n in re.findall(rf"{AWS.upper()}[^\n]*?\s(\d+)\s*$", page.upper(), re.M)
        ],
    }


@pytest.mark.parametrize("case", sorted(CASES))
def test_every_surface_counts_two_per_direction(client, tmp_path, case) -> None:
    make, directions = CASES[case]
    want = 2 * directions
    counts = _counts(client, make(), tmp_path)
    assert counts == {surface: [want] for surface in counts}, counts


def test_the_pdf_counts_sum_to_the_xlsx_total(client, tmp_path) -> None:
    """The flagger sheet's device total carries the 4 signs, same as the XLSX."""
    body = _with(FLAGGER_BODY)
    wb = load_workbook(io.BytesIO(_post(client, "/render/xlsx", body)))
    try:
        rows = [r for r in wb["Device List"].iter_rows(min_row=2, values_only=True) if r[0]]
        total = next(
            r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0] == "Total device count"
        )
    finally:
        wb.close()
    assert total == sum(int(r[5]) for r in rows)
    breakdown = client.post("/render/device-breakdown", json=body, headers=AUTH).json()
    assert breakdown["total_devices"] == total
    assert str(total) in pdf_text(client, body, tmp_path)


def test_a_plan_with_no_direction_count_refuses() -> None:
    """R4: a plan with no direction count refuses rather than guess one
    (Rule 10): a gated kind (lane_closure_divided, whose road facts aren't
    evidenced, Rule 8) and a bare caller with no params."""
    from src.api.schemas import LaneClosureDividedScenario, scenario_to_call

    gated = LaneClosureDividedScenario(
        kind="lane_closure_divided",
        roadType="rural_divided",
        speed=55,
        laneWidth=12.0,
        workType="pavement_repair",
        duration="short",
        workLen=800.0,
        night=False,
        truckMountedAttenuator=False,
        jurisdiction_key="thornton",
    )
    params, generator, kwargs = scenario_to_call(gated)
    assert params.directions_of_travel_affected is None
    placements = generator(params, **kwargs)
    deltas = _jurisdiction_eval(gated, params)[1]
    assert any((d.get("effect") or {}).get("qty_per_direction") for d in deltas)
    for p in (params, None):
        with pytest.raises(ValueError, match="per direction of travel affected"):
            aggregate_device_rows_with_deltas(placements, deltas, p)
        with pytest.raises(ValueError, match="per direction of travel affected"):
            annotate_count_effects(deltas, placements, p)


def test_the_schema_refuses_qty_with_qty_per_direction() -> None:
    import jsonschema

    from src.rules.jurisdiction import _schema

    record = copy.deepcopy(load_jurisdiction("thornton"))
    delta = next(d for d in record["deltas"] if d["effect"].get("qty_per_direction"))
    jsonschema.validate(record, _schema())
    delta["effect"]["qty"] = 4
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(record, _schema())


@pytest.mark.parametrize("key", ["lakewood", "cdot"])
def test_other_advance_warning_rules_keep_their_count(client, key) -> None:
    """R6: Lakewood's and CDOT's advance_warning_signs rules carry no
    per-direction count; whatever they print is unchanged by #388, and
    their deltas carry no new key."""
    for base in (SHOULDER_BODY, FLAGGER_BODY):
        body = _with(base, key=key, street_class="arterial", night=True)
        breakdown = client.post("/render/device-breakdown", json=body, headers=AUTH).json()
        for d in breakdown["jurisdiction"]["applied_deltas"]:
            assert "qty" not in d, d
        assert all(d["qty"] == 1 for d in breakdown["devices"] if str(d["device"]).startswith(AWS))
