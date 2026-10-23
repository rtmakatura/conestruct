"""#374: a flagger plan's AFADs carry one name on every surface.

The layout places an AFAD as a TEMPORARY_SIGNAL labeled AFAD_1/AFAD_2.
Before #374 page 1 and both spreadsheets printed the catalog's "Portable
traffic signal (standard, compact, AFAD, or RDTS)", the crew sheet
"Temporary Traffic Signal", and the breakdown (the screen panel and the
audit's device list) and the PDF legend "Temporary Signal".  Rule 11:
every assertion reads a real render endpoint's output.
"""

from __future__ import annotations

import copy
import io
import json

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.render_api import app
from tests.test_plan_sheet_device_summary import FLAGGER_BODY

AUTH = {"Authorization": "Bearer test-secret"}
NAME = "Automated Flagger Assistance Device (AFAD)"
OLD_NAMES = (
    "Portable traffic signal",
    "Temporary Traffic Signal",
    "Temporary Signal",
)


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setenv("RENDER_API_SECRET", "test-secret")


@pytest.fixture()
def client():
    return TestClient(app)


def _flagger(afad: bool = True) -> dict:
    body = copy.deepcopy(FLAGGER_BODY)
    body["afad"] = afad
    return body


def _post(client: TestClient, path: str, body: dict) -> bytes:
    resp = client.post(path, json=body, headers=AUTH)
    assert resp.status_code == 200, resp.text
    return resp.content


def _pdf_pages(client: TestClient, body: dict, tmp_path) -> list[str]:
    tmp = tmp_path / "sheet.pdf"
    tmp.write_bytes(_post(client, "/render/pdf", body))
    pdf = pdfium.PdfDocument(str(tmp))
    try:
        return [" ".join(page.get_textpage().get_text_range().split()) for page in pdf]
    finally:
        pdf.close()


def _sheet_rows(content: bytes, sheet: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(content))
    try:
        return [r for r in wb[sheet].iter_rows(min_row=2, values_only=True) if r and r[1]]
    finally:
        wb.close()


def test_page_1_names_two_afads(client, tmp_path) -> None:
    page1 = _pdf_pages(client, _flagger(), tmp_path)[0]
    assert f"{NAME} 2" in page1, page1
    assert "Portable traffic signal" not in page1


def test_pdf_legend_names_the_afad(client, tmp_path) -> None:
    pages = _pdf_pages(client, _flagger(), tmp_path)
    assert not any("Temporary Signal" in page for page in pages)
    # Page 1 carries the name twice: the legend row (it fits the box
    # untruncated) and the device summary row.
    assert pages[0].count(NAME) == 2, pages[0]


def test_device_xlsx_names_two_afads(client) -> None:
    rows = _sheet_rows(_post(client, "/render/xlsx", _flagger()), "Device List")
    afads = [r for r in rows if r[1] == "TEMPORARY_SIGNAL"]
    assert len(afads) == 1, afads
    _item, _type, desc, pay_item, unit, qty, _note = afads[0]
    assert (desc, pay_item, unit, qty) == (NAME, "630-86810", "EACH", 2)


def test_quote_xlsx_and_json_name_two_afads(client) -> None:
    rows = _sheet_rows(_post(client, "/render/quote", {"scenario": _flagger()}), "Equipment Detail")
    afads = [r for r in rows if r[1] == "TEMPORARY_SIGNAL"]
    assert len(afads) == 1 and afads[0][3] == NAME and afads[0][4] == 2, afads

    quote = json.loads(_post(client, "/render/quote-breakdown", {"scenario": _flagger()}))
    lines = [ln for ln in quote["equipment_lines"] if ln["device_type"] == "TEMPORARY_SIGNAL"]
    assert [(ln["description"], ln["qty"]) for ln in lines] == [(NAME, 2)]


def test_breakdown_and_audit_device_list_name_two_afads(client) -> None:
    breakdown = json.loads(_post(client, "/render/device-breakdown", _flagger()))
    rows = [r for r in breakdown["devices"] if r["device"] == NAME]
    assert rows == [{"device": NAME, "code": "—", "function": "Traffic control", "qty": 2}]
    assert not any(r["device"] == "Temporary Signal" for r in breakdown["devices"])

    audit = json.loads(_post(client, "/render/audit", {**_flagger(), "include_breakdown": True}))
    assert audit["breakdown"] == breakdown
    assert audit["sections"]["flagger"]["stations_text"] == "2 AFAD"


def test_crew_sheet_names_two_afads(client) -> None:
    md = _post(client, "/render/markdown", _flagger()).decode("utf-8")
    assert md.count(f"2× {NAME}") == 1, md
    assert "Temporary Traffic Signal" not in md


@pytest.mark.parametrize(
    "path,wrap",
    [
        ("/render/xlsx", False),
        ("/render/device-breakdown", False),
        ("/render/quote-breakdown", True),
        ("/render/markdown", False),
    ],
)
def test_no_afad_means_no_afad_name(client, path, wrap) -> None:
    body = _flagger(False)
    content = _post(client, path, {"scenario": body} if wrap else body)
    if path.endswith("xlsx"):
        cells = [str(c) for r in _sheet_rows(content, "Device List") for c in r]
        assert not any(NAME in c for c in cells)
    else:
        assert NAME.encode() not in content
