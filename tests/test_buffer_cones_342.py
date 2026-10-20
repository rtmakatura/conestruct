"""#342: flagger, lane-closure and near-side intersection plans channelize the buffer at <= 2S.

CDOT S-630-1 Sheet 2 General Note 18b: "For tangents along the buffer space
or work area, spacing may not be greater than two times the speed limit"
(MUTCD 11th Ed. Sec 6K.01 para 04, p. 815, the same 2x).  #333 fixed the
shoulder plans; before #342 these three kinds still left the buffer bare
(flagger 45 mph rural, 500 ft zone: the last work-zone cone at 500, the next
device the taper's drum at 860).
"""

from __future__ import annotations

import io
import os
import re
from collections.abc import Iterator
from typing import Any

import pypdfium2 as pdfium
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import TypeAdapter

from src.api import render_api
from src.api.schemas import Scenario, scenario_to_call
from src.generation.layout import (
    buffer_carries_cones,
    buffer_cone_stations,
    generate_flagger_alternating_2lane,
    generate_lane_closure_divided,
)
from src.rules.devices import DeviceType
from src.rules.spacing import buffer_space
from src.rules.validators import ScenarioParams

_TEST_SECRET = "test-secret-do-not-deploy"
AUTH = {"Authorization": f"Bearer {_TEST_SECRET}"}
_ADAPTER: TypeAdapter[Any] = TypeAdapter(Scenario)


@pytest.fixture(scope="module", autouse=True)
def _render_secret() -> Iterator[None]:
    os.environ["RENDER_API_SECRET"] = _TEST_SECRET
    yield


def _params(speed: int, divided: bool, road_type: str = "rural") -> ScenarioParams:
    return ScenarioParams(
        speed_mph=speed,
        num_lanes=2 if divided else 1,
        closure_type="lane",
        road_type=road_type,
        work_zone_length_ft=800.0,
        lane_width_ft=12.0,
        shoulder_width_ft=10.0 if divided else 8.0,
        is_divided=divided,
        jurisdiction="CDOT",
    )


def _ni_body(speed: int = 25, along: float = -60.0, work_len: float = 500.0) -> dict[str, Any]:
    return {
        "kind": "near_intersection",
        "meta": {"project": "T", "address": "", "lat": 0.0, "lng": 0.0, "siteConditions": {}},
        "roadType": "urban_arterial",
        "speed": speed,
        "lanes": 2,
        "laneWidth": 12.0,
        "divided": False,
        "workType": "utility_cut",
        "duration": "short",
        "workLen": work_len,
        "night": False,
        "approaches": [
            {
                "id": "cross_n",
                "speed": 30,
                "roadType": "urban_arterial",
                "lanesPerDirection": 1,
                "laneWidth": 12.0,
                "signalized": False,
                "bearingDeg": None,
                "alongStationFt": along,
            }
        ],
    }


def _flagger_body() -> dict[str, Any]:
    return {
        "kind": "flagger_lane_closure",
        "meta": {"project": "T", "address": "", "lat": 0.0, "lng": 0.0},
        "roadType": "rural_undivided",
        "speed": 45,
        "lanes": 1,
        "laneWidth": 11.0,
        "divided": False,
        "workType": "utility_cut",
        "duration": "short",
        "workLen": 500.0,
        "night": False,
        "pilotCar": False,
        "afad": False,
        "pedestrianAccess": False,
    }


def _ni_placements(body: dict[str, Any]) -> tuple[ScenarioParams, list[Any], Any]:
    params, generator, kwargs = scenario_to_call(_ADAPTER.validate_python(body))
    return params, generator(params, **kwargs), kwargs.get("approaches")


def _assert_no_gap_over_2s(placements: list[Any], wz: float, buf: float, speed: int) -> None:
    # The closed lane's tangent line from the work zone's downstream end to
    # the taper's downstream end: work-zone cones, buffer cones, last drum.
    line = sorted(
        p.station_ft
        for p in placements
        if p.device_type in (DeviceType.CONE, DeviceType.DRUM)
        and p.approach_id == "mainline"
        and 0.0 <= p.station_ft <= wz + buf + 1e-6
    )
    assert line[0] == 0.0 and line[-1] == pytest.approx(wz + buf)
    gaps = [b - a for a, b in zip(line, line[1:], strict=False)]
    assert max(gaps) <= 2 * speed + 1e-6, gaps
    in_buffer = [
        p for p in placements if p.device_type == DeviceType.CONE and wz < p.station_ft < wz + buf
    ]
    assert len(in_buffer) == len(buffer_cone_stations(wz, buf, speed)) > 0


@pytest.mark.parametrize("speed", [25, 35, 45, 55])
def test_flagger_buffer_has_no_gap_over_2s(speed: int) -> None:
    params = _params(speed, divided=False)
    placements = generate_flagger_alternating_2lane(params)
    buf = buffer_space(speed, jurisdiction="CDOT")
    _assert_no_gap_over_2s(placements, params.work_zone_length_ft, buf, speed)
    # On the centerline, where the one-lane two-way taper ends.
    wz = params.work_zone_length_ft
    assert all(
        p.offset_ft == 0.0
        for p in placements
        if p.device_type == DeviceType.CONE and wz < p.station_ft < wz + buf
    )


@pytest.mark.parametrize("speed", [45, 55, 65])
def test_lane_closure_buffer_has_no_gap_over_2s(speed: int) -> None:
    params = _params(speed, divided=True)
    placements = generate_lane_closure_divided(params)
    _assert_no_gap_over_2s(
        placements, params.work_zone_length_ft, buffer_space(speed, jurisdiction="CDOT"), speed
    )


@pytest.mark.parametrize("speed", [25, 35, 45])
def test_near_side_buffer_has_no_gap_over_2s(speed: int) -> None:
    params, placements, approaches = _ni_placements(_ni_body(speed=speed))
    assert buffer_carries_cones(params, approaches)
    buf = buffer_space(
        speed, jurisdiction=params.jurisdiction, work_zone_speed_mph=params.work_zone_speed_mph
    )
    _assert_no_gap_over_2s(placements, params.work_zone_length_ft, buf, speed)


def test_the_far_side_keeps_its_companion_line_only() -> None:
    # Cross street past the work (along > workLen): the §6N.12.12 companion
    # tangent already delineates the buffer, so no buffer cones are added.
    params, placements, approaches = _ni_placements(_ni_body(speed=25, along=900.0))
    assert not buffer_carries_cones(params, approaches)


def _surfaces(body: dict[str, Any]) -> tuple[dict[str, Any], int, str, str]:
    c = TestClient(render_api.app)
    sp = c.post("/render/audit", headers=AUTH, json=body).json()["sections"]["spacing"]
    xl = load_workbook(io.BytesIO(c.post("/render/xlsx", headers=AUTH, json=body).content))
    xl_cones = sum(
        r[5]
        for r in xl["Device List"].iter_rows(min_row=2, values_only=True)
        if r[2] and "Cone" in str(r[2])
    )
    md = c.post("/render/markdown", headers=AUTH, json=body).text
    pdf = pdfium.PdfDocument(c.post("/render/pdf", headers=AUTH, json=body).content)
    return sp, xl_cones, md, pdf[0].get_textpage().get_text_range()


def test_every_flagger_surface_counts_and_names_the_buffer_cones() -> None:
    # 360 ft buffer at 45 mph: ceil(360 / 90) = 4 intervals of 90 ft -> 3 cones.
    sp, xl_cones, md, page1 = _surfaces(_flagger_body())
    assert sp["n_buffer_cones"] == 3
    assert sp["buffer_count_text"].startswith(
        "Buffer: 360 ft / 90 ft max spacing -> 3 cones at 90.0 ft intervals"
    )
    total = sp["n_cones_placed_total"]
    assert " + 3 buffer + " in sp["downstream_count_text"]
    assert xl_cones == total
    assert "Then continue the line through the buffer space: 3 more traffic cones at 90 ft" in md
    assert "BUFFER = 360 ft (NTS) · 3 CONES" in page1
    summary = re.search(r"Traffic Cone \(\d+-inch\) (\d+)", page1)
    assert summary is not None and int(summary.group(1)) == total


def test_every_near_side_surface_counts_and_names_the_buffer_cones() -> None:
    # 155 ft buffer at 25 mph: ceil(155 / 50) = 4 intervals of 38.75 ft -> 3 cones.
    sp, xl_cones, md, page1 = _surfaces(_ni_body())
    assert sp["n_buffer_cones"] == 3
    assert sp["buffer_count_text"].startswith(
        "Buffer: 155 ft / 50 ft max spacing -> 3 cones at 38.8 ft intervals"
    )
    assert xl_cones == sp["n_cones_placed_total"]
    assert "Then continue the line through the buffer space: 3 more traffic cones at 39 ft" in md
    # Not described as the far side's companion closure.
    assert "The lane must already be closed where it crosses the intersection" not in md
    assert "BUFFER = 155 ft (NTS) · 3 CONES" in page1
