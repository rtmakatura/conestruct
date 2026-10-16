"""#324: a saved plan with no quote settings of its own (a row saved before
#202's column existed, ``quote_settings`` null) prices its quote at the
defaults, and the XLSX says so (Rule 10).

The site's row route sends ``settings_source: "default"`` for such a row
(``lib/render-proxy.ts`` ``proxyRender``; pinned in
``app/api/plans/quote-settings.test.ts``).  Here, the payload level: the line
lands at Quote Summary A8 only with that flag, and every other cell of every
sheet, the prices included, is the same with or without it.
"""

from __future__ import annotations

import io
import os
import re
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.render_api import app
from src.export.quote_generator import QUOTE_DEFAULT_RATES_LINE
from tests.test_quote_settings import _shoulder_scenario

_STAMP = re.compile(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")

DEFAULTS = {
    "project_duration_days": 1,
    "num_flaggers": 0,
    "delivery_distance_miles": 20.0,
    "overhead_pct": 0.10,
    "profit_pct": 0.10,
    "flagger_hourly_rate": 55.0,
    "tcs_hourly_rate": 75.0,
    "crew_hourly_rate": 45.0,
}


@pytest.fixture(scope="module", autouse=True)
def _secret() -> None:
    os.environ["RENDER_API_SECRET"] = "test-secret-do-not-deploy"


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app)


def quote(client: TestClient, **extra: Any) -> dict[str, dict[str, Any]]:
    r = client.post(
        "/render/quote",
        headers={"Authorization": "Bearer test-secret-do-not-deploy"},
        json={"scenario": _shoulder_scenario(), "settings": DEFAULTS, **extra},
    )
    assert r.status_code == 200, r.text[:300]
    wb = load_workbook(io.BytesIO(r.content))
    # The "Generated" timestamp ticks between two requests that straddle a
    # second, so it's left out of the comparison.
    return {
        ws.title: {
            c.coordinate: c.value
            for row in ws.iter_rows()
            for c in row
            if c.value is not None and not _STAMP.match(str(c.value))
        }
        for ws in wb.worksheets
    }


def test_a_plan_with_no_saved_settings_says_its_rates_are_the_defaults(
    client: TestClient,
) -> None:
    cells = quote(client, settings_source="default")
    assert cells["Quote Summary"]["A8"] == QUOTE_DEFAULT_RATES_LINE


def test_without_the_flag_the_quote_is_unchanged(client: TestClient) -> None:
    plain = quote(client)
    flagged = quote(client, settings_source="default")
    assert "A8" not in plain["Quote Summary"]
    # Every cell but the new line is the same, the prices included.
    flagged["Quote Summary"].pop("A8")
    assert flagged == plain


def test_an_unknown_source_is_refused(client: TestClient) -> None:
    r = client.post(
        "/render/quote",
        headers={"Authorization": "Bearer test-secret-do-not-deploy"},
        json={"scenario": _shoulder_scenario(), "settings": DEFAULTS, "settings_source": "saved"},
    )
    assert r.status_code == 422
